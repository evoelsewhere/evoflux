"""End-to-end check of a built LibreOffice runtime bundle on this machine.

Installs the bundle through the sidecar's real installer (download, size and
SHA-256 check, ed25519 signature, guarded extraction, activation) into a
throwaway data directory, then renders generated DOCX, PPTX and XLSX files
through the document preview pipeline and asserts that LibreOffice produced
them with the expected text. CI runs this on every runtime platform so macOS
and Windows bundles are proven before they are published.

Example::

    uv run --extra office-preview python scripts/verify_office_runtime.py \\
        --manifest dist/office-runtime/libreoffice-26.8.0-win32-x64.json \\
        --public-key keys/office.public.pem --key-id evoflux-office-1
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

MARKERS = {
    ".docx": "Quarterly revenue review",
    ".pptx": "Launch readiness",
    ".xlsx": "Forecast total",
}


def _log(message: str) -> None:
    print(message, flush=True)


def _fixtures(directory: Path) -> list[Path]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    document = Document()
    document.add_heading(MARKERS[".docx"], level=1)
    document.add_paragraph("Revenue grew in every region.", style="List Number")
    document.add_paragraph("Margins held steady.", style="List Number")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Growth"
    docx = directory / "report.docx"
    document.save(docx)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = MARKERS[".pptx"]
    slide.placeholders[1].text = "Checklist complete"
    presentation.slides.add_slide(
        presentation.slide_layouts[5]
    ).shapes.title.text = "Q&A"
    pptx = directory / "deck.pptx"
    presentation.save(pptx)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forecast"
    sheet.append(["Month", "Units"])
    for month, units in (("Jan", 120), ("Feb", 135), ("Mar", 150)):
        sheet.append([month, units])
    sheet.append([MARKERS[".xlsx"], "=SUM(B2:B4)"])
    # Wide enough that the label is not clipped by the value beside it, as
    # it would be in Excel too.
    sheet.column_dimensions["A"].width = 24
    workbook.create_sheet("Notes")["A1"] = "Assumptions"
    xlsx = directory / "model.xlsx"
    workbook.save(xlsx)
    return [docx, pptx, xlsx]


async def _install() -> None:
    from app.services.office_runtime import installer

    status = installer.runtime_status()
    if not status.available:
        raise SystemExit(
            "The manifest did not yield a verified asset for this platform"
        )
    installer.start_runtime_install()
    while True:
        await asyncio.sleep(1)
        job = installer.runtime_status().job
        if job is None:
            break
        if job.phase == "failed":
            raise SystemExit(f"Install failed: {job.error}")
    runtime = installer.installed_runtime()
    if runtime is None:
        raise SystemExit("Install finished without an active runtime")
    _log(f"installed LibreOffice {runtime.version} at {runtime.executable}")


def _render(fixtures: list[Path]) -> None:
    from app.services.document_preview.service import _render_source

    expected_pages = {".docx": 1, ".pptx": 2, ".xlsx": 2}
    for fixture in fixtures:
        started = time.perf_counter()
        rendered = _render_source(fixture)
        elapsed = time.perf_counter() - started
        engine = re.search(r'data-preview-renderer="([^"]+)"', rendered)
        if engine is None or not engine.group(1).startswith("libreoffice-"):
            raise SystemExit(f"{fixture.name} was not rendered by LibreOffice")
        pages = rendered.count("data-preview-item")
        if pages < expected_pages[fixture.suffix]:
            raise SystemExit(f"{fixture.name}: expected pages, got {pages}")
        if MARKERS[fixture.suffix] not in rendered:
            raise SystemExit(f"{fixture.name}: text layer is missing its marker")
        _log(
            f"ok {fixture.name}: {pages} page(s) via {engine.group(1)} in {elapsed:.1f}s"
        )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n", 1)[0])
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--public-key", type=Path, required=True)
    parser.add_argument("--key-id", required=True)
    parser.add_argument(
        "--archive",
        type=Path,
        help="local archive to install instead of the manifest URL (not yet published)",
    )
    args = parser.parse_args()

    assets = json.loads(args.manifest.read_text(encoding="utf-8"))
    if args.archive is not None:
        # The signature covers content and identity, not location, so a
        # local copy verifies exactly like the published asset.
        for entry in assets.values():
            entry["url"] = args.archive.resolve().as_uri()
    with tempfile.TemporaryDirectory(prefix="evoflux-lo-verify-") as directory:
        root = Path(directory)
        override = root / "manifest.json"
        override.write_text(
            json.dumps(
                {
                    "keys": {args.key_id: args.public_key.read_text(encoding="ascii")},
                    "assets": assets,
                }
            ),
            encoding="utf-8",
        )
        os.environ["EVOFLUX_OFFICE_RUNTIME_MANIFEST"] = str(override)
        os.environ["EVOFLUX_DATA_DIR"] = str(root / "data")
        os.environ["EVOFLUX_CACHE_DIR"] = str(root / "cache")
        fixtures_dir = root / "fixtures"
        fixtures_dir.mkdir()
        fixtures = _fixtures(fixtures_dir)
        asyncio.run(_install())
        _render(fixtures)
    _log("office runtime verified")


if __name__ == "__main__":
    main()

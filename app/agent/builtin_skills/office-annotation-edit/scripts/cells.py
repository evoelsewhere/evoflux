#!/usr/bin/env python3
"""Describe the cells an EvoFlux workbook annotation points at.

Given a workbook, a sheet name and the annotation's A1 range, print JSON
describing each cell: its coordinate, what it holds (a formula or a literal),
the last computed value Excel or LibreOffice saved (``null`` until the
workbook is recalculated), number format, font, fill, and the merged block it
belongs to. Empty cells in the range are skipped.

Usage:
    python cells.py model.xlsx --sheet Model --range B3:D8
    python cells.py model.xlsx --sheet "Q3 Plan" --range C12

Run with ``uv run --with openpyxl python cells.py …``.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

# A range this large is a whole-sheet selection; describe its first cells only.
MAX_CELLS = 400


def _colour(colour) -> str | None:
    value = getattr(colour, "rgb", None)
    return value if isinstance(value, str) and len(value) in (6, 8) else None


def describe(cell, cached, merged: dict[str, str]) -> dict:
    holds = cell.value
    entry: dict = {"cell": cell.coordinate}
    if isinstance(holds, str) and holds.startswith("="):
        entry["formula"] = holds
        entry["value"] = cached
    else:
        entry["value"] = holds if not hasattr(holds, "isoformat") else holds.isoformat()
    if cell.number_format and cell.number_format != "General":
        entry["number_format"] = cell.number_format
    font = {
        key: value
        for key, value in {
            "bold": bool(cell.font.b) or None,
            "italic": bool(cell.font.i) or None,
            "size": cell.font.sz if cell.font.sz not in (None, 11) else None,
            "color": _colour(cell.font.color),
        }.items()
        if value
    }
    if font:
        entry["font"] = font
    fill = _colour(cell.fill.fgColor) if cell.fill and cell.fill.fill_type == "solid" else None
    if fill:
        entry["fill"] = fill
    if cell.coordinate in merged:
        entry["merged"] = merged[cell.coordinate]
    return entry


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Describe the cells of an annotated range.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", required=True, help="sheet name from the annotation")
    parser.add_argument("--range", required=True, dest="cells", help='A1 range, e.g. "B3:D8"')
    args = parser.parse_args(argv)

    workbook = load_workbook(args.workbook)
    computed = load_workbook(args.workbook, data_only=True)
    if args.sheet not in workbook.sheetnames:
        print(f"error: no sheet {args.sheet!r}; sheets: {workbook.sheetnames}", file=sys.stderr)
        return 1
    sheet, cached_sheet = workbook[args.sheet], computed[args.sheet]
    try:
        min_col, min_row, max_col, max_row = range_boundaries(args.cells.replace("$", ""))
    except (TypeError, ValueError):
        print(f"error: {args.cells!r} is not an A1 range", file=sys.stderr)
        return 1
    merged = {
        coordinate: str(block)
        for block in sheet.merged_cells.ranges
        for row in sheet.iter_rows(
            min_row=block.min_row, max_row=block.max_row, min_col=block.min_col, max_col=block.max_col
        )
        for coordinate in (cell.coordinate for cell in row)
    }
    cells = []
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if cell.value is None and cell.coordinate not in merged:
                continue
            cells.append(describe(cell, cached_sheet[cell.coordinate].value, merged))
            if len(cells) >= MAX_CELLS:
                break
        if len(cells) >= MAX_CELLS:
            break
    print(json.dumps(
        {"sheet": sheet.title, "range": args.cells, "cells": cells, "truncated": len(cells) >= MAX_CELLS},
        ensure_ascii=False,
        indent=2,
        default=str,
    ))
    return 0


if __name__ == "__main__":
    sys.exit(main())

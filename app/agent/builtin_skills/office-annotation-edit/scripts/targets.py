#!/usr/bin/env python3
"""Describe the shapes an EvoFlux annotation points at.

Given a deck, a 1-based slide number and the annotation's shape ids and/or
selected area (percent of the slide), print JSON describing each target:
id, name, kind, box (percent and EMU), text, solid fills and, for charts, the
chart type with its series and their colours. With only an area, every shape
the box covers is listed, most-covered first, with the share of the shape
inside the box.

Usage:
    python targets.py deck.pptx --slide 7 --shape 12
    python targets.py deck.pptx --slide 7 --area 8.1,21.5,84,52.3
    python targets.py deck.pptx --slide 7 --shape 12 --shape 14 --area 8,20,80,50

Groups are searched recursively; a shape inside a group reports the group's
id as ``group``. Run with ``uv run --with python-pptx python targets.py …``.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _parse_area(value: str) -> tuple[float, float, float, float]:
    try:
        x, y, w, h = (float(part) for part in value.split(","))
    except ValueError as exc:
        raise argparse.ArgumentTypeError("area is x,y,w,h in percent") from exc
    return x, y, w, h


def _walk(shapes, group=None):
    for shape in shapes:
        yield shape, group
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes, group=shape.shape_id)


def _box(shape, width: int, height: int) -> dict | None:
    if shape.left is None or shape.top is None or not shape.width or not shape.height:
        return None
    return {
        "x": round(100 * shape.left / width, 2),
        "y": round(100 * shape.top / height, 2),
        "w": round(100 * shape.width / width, 2),
        "h": round(100 * shape.height / height, 2),
        "emu": [int(shape.left), int(shape.top), int(shape.width), int(shape.height)],
    }


def _overlap(box: dict, area: tuple[float, float, float, float]) -> float:
    ax, ay, aw, ah = area
    left, top = max(box["x"], ax), max(box["y"], ay)
    right = min(box["x"] + box["w"], ax + aw)
    bottom = min(box["y"] + box["h"], ay + ah)
    if right <= left or bottom <= top:
        return 0.0
    return (right - left) * (bottom - top) / max(box["w"] * box["h"], 1e-9)


def _solid_fill(fill) -> str | None:
    try:
        if (
            fill.type == 1 and fill.fore_color and fill.fore_color.type is not None
        ):  # SOLID
            color = fill.fore_color
            return f"#{color.rgb}" if color.type == 1 else f"theme:{color.theme_color}"
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _kind(shape) -> str:
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(shape, "is_placeholder", False) and shape.is_placeholder:
        return "placeholder"
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return "text"
    return "shape"


def _chart(shape) -> dict:
    chart = shape.chart
    series = []
    for plot in chart.plots:
        for item in plot.series:
            entry = {"name": item.name, "values": list(item.values)[:24]}
            try:
                entry["fill"] = _solid_fill(item.format.fill)
            except (AttributeError, ValueError):
                entry["fill"] = None
            series.append(entry)
    return {"type": str(chart.chart_type), "series": series}


def describe(shape, group, width: int, height: int) -> dict:
    info: dict = {
        "id": shape.shape_id,
        "name": shape.name,
        "kind": _kind(shape),
        "box": _box(shape, width, height),
    }
    if group is not None:
        info["group"] = group
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        info["text"] = shape.text_frame.text[:500]
    if getattr(shape, "is_placeholder", False) and shape.is_placeholder:
        info["placeholder"] = {
            "idx": shape.placeholder_format.idx,
            "type": str(shape.placeholder_format.type),
        }
    try:
        fill = _solid_fill(shape.fill)
        if fill:
            info["fill"] = fill
    except (AttributeError, TypeError):
        pass
    if info["kind"] == "chart":
        info["chart"] = _chart(shape)
    if info["kind"] == "table":
        table = shape.table
        info["table"] = {"rows": len(table.rows), "columns": len(table.columns)}
    return info


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Describe the shapes an EvoFlux annotation points at."
    )
    parser.add_argument("deck", type=Path)
    parser.add_argument("--slide", type=int, required=True, help="1-based slide number")
    parser.add_argument(
        "--shape", type=int, action="append", default=[], help="shape id"
    )
    parser.add_argument("--area", type=_parse_area, help="x,y,w,h in percent")
    args = parser.parse_args(argv)
    if not args.shape and args.area is None:
        parser.error("give --shape and/or --area")

    prs = Presentation(str(args.deck))
    if not 1 <= args.slide <= len(prs.slides):
        print(
            f"slide {args.slide} not found; the deck has {len(prs.slides)}",
            file=sys.stderr,
        )
        return 2
    width, height = int(prs.slide_width or 0), int(prs.slide_height or 0)
    if width <= 0 or height <= 0:
        print("the deck has no slide size", file=sys.stderr)
        return 2
    slide = prs.slides[args.slide - 1]
    shapes = list(_walk(slide.shapes))

    result: dict = {
        "slide": args.slide,
        "slide_size_emu": [width, height],
        "targets": [],
    }
    by_id = {shape.shape_id: (shape, group) for shape, group in shapes}
    missing = [shape_id for shape_id in args.shape if shape_id not in by_id]
    for shape_id in args.shape:
        if shape_id in by_id:
            result["targets"].append(describe(*by_id[shape_id], width, height))
    if args.area is not None:
        covered = []
        for shape, group in shapes:
            if shape.shape_id in args.shape:
                continue
            box = _box(shape, width, height)
            share = _overlap(box, args.area) if box else 0.0
            if share > 0.05:
                info = describe(shape, group, width, height)
                info["covered"] = round(share, 2)
                covered.append(info)
        covered.sort(key=lambda item: item["covered"], reverse=True)
        result["in_area"] = covered
    if missing:
        result["missing_shape_ids"] = missing
    json.dump(result, sys.stdout, ensure_ascii=False, indent=1)
    print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Build a deck slide by slide so the user watches it take shape.

EvoFlux opens a live preview as soon as the deck file appears and redraws it
every time the file is saved: finished slides show, and the rest appear as
placeholders titled from the plan. So create the file first, then add one
slide per command, writing the next slide while the user looks at the last.
A single script that builds every slide finishes in about a second, and the
user only ever sees the finished deck.

Commands:

    init   DECK --title "Cover" --title "Agenda" ...  create an empty deck
           [--placeholder cover ...]                  carrying the plan; the
                                                      optional placeholders
                                                      (one per title) only
                                                      shape the grey preview
                                                      shown until each slide
                                                      exists
    add    DECK slides/02_agenda.py [--replace N]     run the file's
                                                      ``build(prs)`` to add
                                                      exactly one slide, then
                                                      save (``--replace N``
                                                      rebuilds slide N)
    mark   DECK                                       re-embed the plan after
                                                      a generator rewrote the
                                                      file (PptxGenJS)
    finish DECK                                       mark the deck complete

A slide file defines ``build(prs)`` and adds one slide to ``prs``. Its folder
is importable, so shared palette and helpers can live beside it (for example
``slides/theme.py``, imported with ``from theme import ...``).

Saves are atomic (temporary file + rename), so the preview never reads a
half-written deck. The plan is stored in the ``evoflux.deck`` custom document
property and in a scratch copy under the system temp directory, which `mark`
uses when a generator rewrites the file from scratch.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

PLAN_PROPERTY = "evoflux.deck"
CUSTOM_PART = "docProps/custom.xml"
CUSTOM_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
CUSTOM_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
)
_FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"
# Shapes the preview can draw for a slide that does not exist yet. They only
# decorate the loading state; the slide itself is designed from the outline.
# Keep in step with ``PLACEHOLDER_KINDS`` in
# app/services/document_preview/live_deck.py.
PLACEHOLDERS = (
    "cover",
    "bullets",
    "split",
    "cards",
    "chart",
    "line",
    "donut",
    "stats",
    "table",
    "timeline",
    "diagram",
    "quote",
    "closing",
)


class DeckLiveError(RuntimeError):
    """The deck or its plan is missing or malformed."""


def _plan_path(deck: Path) -> Path:
    """Scratch copy of the plan, outside the workspace so it never shows up
    among the user's files; keyed by the deck's absolute path."""
    key = hashlib.sha256(str(deck.resolve()).encode("utf-8")).hexdigest()[:24]
    directory = Path(tempfile.gettempdir()) / "evoflux-deck-live"
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{key}.json"


def _custom_xml(plan: dict) -> bytes:
    value = escape(json.dumps(plan, ensure_ascii=False, separators=(",", ":")))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        f'<property fmtid="{_FMTID}" pid="2" name="{PLAN_PROPERTY}">'
        f"<vt:lpwstr>{value}</vt:lpwstr></property></Properties>"
    ).encode("utf-8")


def _with_plan(package: bytes, plan: dict) -> bytes:
    """Return ``package`` with the plan written into its custom properties."""
    source = zipfile.ZipFile(io.BytesIO(package))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename == CUSTOM_PART:
                continue  # rewritten below; only EvoFlux writes this part here
            payload = source.read(info)
            if info.filename == "[Content_Types].xml" and CUSTOM_PART.encode() not in payload:
                payload = payload.replace(
                    b"</Types>",
                    f'<Override PartName="/{CUSTOM_PART}" '
                    f'ContentType="{CUSTOM_CONTENT_TYPE}"/></Types>'.encode(),
                )
            elif info.filename == "_rels/.rels" and CUSTOM_REL_TYPE.encode() not in payload:
                payload = payload.replace(
                    b"</Relationships>",
                    f'<Relationship Id="rIdEvofluxDeck" Type="{CUSTOM_REL_TYPE}" '
                    f'Target="{CUSTOM_PART}"/></Relationships>'.encode(),
                )
            target.writestr(info, payload)
        target.writestr(CUSTOM_PART, _custom_xml(plan))
    return output.getvalue()


def _atomic_write(deck: Path, payload: bytes) -> None:
    temporary = deck.with_name(f".{deck.name}.{os.getpid()}.tmp")
    temporary.write_bytes(payload)
    for attempt in range(40):
        try:
            os.replace(temporary, deck)
            return
        except PermissionError:
            # Windows: the previewer may be reading the deck for a moment.
            if attempt == 39:
                temporary.unlink(missing_ok=True)
                raise
            time.sleep(0.1)


def _read_plan(deck: Path) -> dict:
    try:
        return json.loads(_plan_path(deck).read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise DeckLiveError(f"No live plan for {deck.name}; run `init` first.") from exc


def _write(deck: Path, package: bytes, plan: dict) -> None:
    _plan_path(deck).write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    _atomic_write(deck, _with_plan(package, plan))


def init(deck: Path, titles: list[str], placeholders: list[str] | None = None) -> dict:
    """Create an empty 16:9 deck that carries the build plan.

    ``placeholders`` picks, per title, the rough shape (see ``PLACEHOLDERS``)
    of the grey stand-in the preview shows until that slide is added. It has
    no effect on the deck. Without it the first stand-in is a cover, the last
    a closing slide and the rest vary.
    """
    from pptx import Presentation
    from pptx.util import Inches

    if not titles:
        raise DeckLiveError("Give one --title per planned slide.")
    placeholders = placeholders or []
    if placeholders and len(placeholders) != len(titles):
        raise DeckLiveError(
            f"Give one --placeholder per --title, or none ({len(titles)} titles, "
            f"{len(placeholders)} placeholders)."
        )
    unknown = sorted(set(placeholders) - set(PLACEHOLDERS))
    if unknown:
        raise DeckLiveError(
            f"Unknown placeholder {', '.join(unknown)}; choose from {', '.join(PLACEHOLDERS)}."
        )
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    buffer = io.BytesIO()
    presentation.save(buffer)
    plan = {"v": 1, "total": len(titles), "titles": titles, "state": "building"}
    if placeholders:
        plan["placeholders"] = placeholders
    # EvoFlux names the session running this command; only that session's
    # viewer shows the build live, and only while its turn runs.
    session = os.environ.get("EVOFLUX_SESSION", "").strip()
    if session:
        plan["session"] = session
    deck.parent.mkdir(parents=True, exist_ok=True)
    _write(deck, buffer.getvalue(), plan)
    return plan


def mark(deck: Path, *, done: bool = False) -> dict:
    """Re-embed the plan into ``deck`` as it is now (e.g. after PptxGenJS)."""
    plan = _read_plan(deck)
    if done:
        plan["state"] = "done"
    try:
        package = deck.read_bytes()
    except OSError as exc:
        raise DeckLiveError(f"Cannot read {deck}") from exc
    _write(deck, package, plan)
    if done:
        _plan_path(deck).unlink(missing_ok=True)
    return plan


def _load_builder(script: Path):
    """Return the ``build`` function a slide file defines."""
    import importlib.util

    if not script.is_file():
        raise DeckLiveError(f"No slide file at {script}")
    folder = str(script.resolve().parent)
    if folder not in sys.path:
        sys.path.insert(0, folder)  # shared helpers beside the slide files
    spec = importlib.util.spec_from_file_location(f"_evoflux_slide_{script.stem}", script)
    if spec is None or spec.loader is None:
        raise DeckLiveError(f"Cannot load {script}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    build = getattr(module, "build", None)
    if not callable(build):
        raise DeckLiveError(f"{script.name} must define build(prs)")
    return build


def _move_slide(presentation, old: int, new: int) -> None:
    slides = presentation.slides._sldIdLst
    entry = list(slides)[old]
    slides.remove(entry)
    slides.insert(new, entry)


def _drop_slide(presentation, index: int) -> None:
    slides = presentation.slides._sldIdLst
    entry = list(slides)[index]
    presentation.part.drop_rel(entry.rId)
    slides.remove(entry)


def add(deck: Path, script: Path, *, replace: int | None = None) -> int:
    """Add the one slide ``script`` builds (or rebuild slide ``replace``).

    The deck is saved only when the build succeeds, so a failing slide file
    leaves the previous state on screen. Returns the new slide count.
    """
    from pptx import Presentation

    plan = _read_plan(deck)
    presentation = Presentation(str(deck))
    before = len(presentation.slides)
    if replace is not None and not 1 <= replace <= before:
        raise DeckLiveError(f"--replace {replace}: the deck has {before} slides")
    _load_builder(script)(presentation)
    added = len(presentation.slides) - before
    if added != 1:
        raise DeckLiveError(f"{script.name} must add exactly one slide (it added {added})")
    if replace is not None:
        _move_slide(presentation, before, replace - 1)
        _drop_slide(presentation, replace)  # the old slide, now one further on
    buffer = io.BytesIO()
    presentation.save(buffer)
    _write(deck, buffer.getvalue(), plan)
    return len(presentation.slides)


class LiveDeck:
    """python-pptx helper: save after every slide, finish at the end.

    Prefer ``add``: a script that builds every slide in one run finishes in
    about a second, so the user never sees the deck take shape.
    """

    def __init__(self, deck: str | os.PathLike[str]) -> None:
        self.path = Path(deck)
        self.plan = _read_plan(self.path)

    def open(self):
        from pptx import Presentation

        return Presentation(str(self.path))

    def save(self, presentation) -> None:
        buffer = io.BytesIO()
        presentation.save(buffer)
        _write(self.path, buffer.getvalue(), self.plan)

    def finish(self, presentation) -> None:
        self.plan["state"] = "done"
        self.save(presentation)
        _plan_path(self.path).unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n", 1)[0])
    commands = parser.add_subparsers(dest="command", required=True)
    start = commands.add_parser("init", help="create an empty deck carrying the plan")
    start.add_argument("deck", type=Path)
    start.add_argument("--title", action="append", default=[], help="one per planned slide")
    start.add_argument(
        "--placeholder",
        action="append",
        default=[],
        choices=PLACEHOLDERS,
        help="optional, one per --title in order: the loading preview's rough shape "
        "(cosmetic; does not affect the slide)",
    )
    one = commands.add_parser("add", help="add the one slide a slide file's build(prs) makes")
    one.add_argument("deck", type=Path)
    one.add_argument("script", type=Path)
    one.add_argument("--replace", type=int, metavar="N", help="rebuild slide N instead")
    again = commands.add_parser("mark", help="re-embed the plan after a generator rewrote the deck")
    again.add_argument("deck", type=Path)
    end = commands.add_parser("finish", help="mark the deck complete")
    end.add_argument("deck", type=Path)
    args = parser.parse_args()
    try:
        if args.command == "init":
            plan = init(args.deck, args.title, args.placeholder)
            print(f"created {args.deck} with {plan['total']} planned slides")
        elif args.command == "add":
            count = add(args.deck, args.script, replace=args.replace)
            total = _read_plan(args.deck)["total"]
            if args.replace is not None:
                print(f"rebuilt slide {args.replace} of {args.deck}")
            elif count < total:
                print(f"slide {count} of {total} saved; write the next slide file")
            else:
                print(f"slide {count} of {total} saved; run `finish` when the deck is done")
        elif args.command == "mark":
            mark(args.deck)
            print(f"re-embedded the plan in {args.deck}")
        else:
            mark(args.deck, done=True)
            print(f"{args.deck} is complete")
    except DeckLiveError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

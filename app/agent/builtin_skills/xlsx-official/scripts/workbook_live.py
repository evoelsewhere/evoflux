#!/usr/bin/env python3
"""Build a workbook sheet by sheet so the user watches it take shape.

EvoFlux opens a live preview as soon as the workbook file appears and redraws
it every time the file is saved: finished sheets show, and the rest appear as
loading skeletons. So create the file first, then add one sheet per command,
writing the next sheet while the user looks at the last. A single script that
builds every sheet finishes in about a second, and the user only ever sees
the finished workbook.

Commands:

    init    BOOK --sheets N                       create a workbook expecting
                                                  N sheets
    add     BOOK sheets/02_data.py [--replace N]  run the file's ``build(wb)``
                                                  to add exactly one sheet,
                                                  then save (``--replace N``
                                                  rebuilds sheet N in place)
    finish  BOOK                                  mark the workbook complete
    rebuild BOOK sheets/                          rebuild every sheet from the
                                                  numbered sheet files in one
                                                  save

After `finish`, fix a sheet with `add BOOK sheets/03_x.py --replace 3`, or run
`rebuild` after changing shared helpers or several sheets: the workbook stays
finished and the preview simply redraws. Never run `init` again on a workbook
that has sheets; it empties it and the preview starts over from nothing.

A sheet file defines ``build(wb)``, which adds one worksheet to the openpyxl
``Workbook`` (``wb.create_sheet("Data")``) and fills it; formulas may refer to
sheets added before it. Its folder is importable, so shared styles and
helpers can live beside it (``sheets/theme.py``, ``from theme import ...``).

Saves are atomic (temporary file + rename), so the preview never reads a
half-written workbook. The plan lives in the ``evoflux.deck`` custom document
property and in a scratch copy under the system temp directory. The zip and
plan helpers are kept in step with ``pptx-official/scripts/deck_live.py``.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib
import io
import json
import os
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

PLAN_PROPERTY = "evoflux.deck"
CUSTOM_PART = "docProps/custom.xml"
CUSTOM_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
CUSTOM_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
)
_FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"
_MAX_SHEETS = 500
# The stand-in sheet a new workbook holds until its first real sheet lands.
_STAND_IN = "Building…"


class WorkbookLiveError(RuntimeError):
    """The workbook or its plan is missing or malformed."""


def _plan_path(book: Path) -> Path:
    """Scratch copy of the plan, outside the workspace so it never shows up
    among the user's files; keyed by the workbook's absolute path."""
    key = hashlib.sha256(str(book.resolve()).encode("utf-8")).hexdigest()[:24]
    directory = Path(tempfile.gettempdir()) / "evoflux-deck-live"
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{key}.json"


def _custom_xml(plan: dict) -> bytes:
    value = escape(json.dumps(plan, ensure_ascii=False, separators=(",", ":")))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        f'<property fmtid="{_FMTID}" pid="2" name="{PLAN_PROPERTY}">'
        f"<vt:lpwstr>{value}</vt:lpwstr></property></Properties>"
    ).encode("utf-8")


def _with_plan(package: bytes, plan: dict) -> bytes:
    """Return ``package`` with the plan written into its custom properties."""
    source = zipfile.ZipFile(io.BytesIO(package))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename == CUSTOM_PART:
                continue  # rewritten below; only EvoFlux writes this part here
            payload = source.read(info)
            if info.filename == "[Content_Types].xml" and CUSTOM_PART.encode() not in payload:
                payload = payload.replace(
                    b"</Types>",
                    f'<Override PartName="/{CUSTOM_PART}" '
                    f'ContentType="{CUSTOM_CONTENT_TYPE}"/></Types>'.encode(),
                )
            elif info.filename == "_rels/.rels" and CUSTOM_REL_TYPE.encode() not in payload:
                payload = payload.replace(
                    b"</Relationships>",
                    f'<Relationship Id="rIdEvofluxDeck" Type="{CUSTOM_REL_TYPE}" '
                    f'Target="{CUSTOM_PART}"/></Relationships>'.encode(),
                )
            target.writestr(info, payload)
        target.writestr(CUSTOM_PART, _custom_xml(plan))
    return output.getvalue()


def _atomic_write(book: Path, payload: bytes) -> None:
    temporary = book.with_name(f".{book.name}.{os.getpid()}.tmp")
    temporary.write_bytes(payload)
    for attempt in range(40):
        try:
            os.replace(temporary, book)
            return
        except PermissionError:
            # Windows: the previewer may be reading the workbook for a moment.
            if attempt == 39:
                temporary.unlink(missing_ok=True)
                raise
            time.sleep(0.1)


def _embedded_plan(book: Path) -> dict | None:
    """The plan stored in the workbook itself, if any."""
    import xml.etree.ElementTree as ElementTree

    try:
        with zipfile.ZipFile(book) as archive:
            info = archive.getinfo(CUSTOM_PART)
            if info.file_size > 64 * 1024:
                return None
            root = ElementTree.fromstring(archive.read(info))
    except (OSError, KeyError, zipfile.BadZipFile, ElementTree.ParseError):
        return None
    for prop in root.iter():
        if prop.get("name") == PLAN_PROPERTY:
            try:
                plan = json.loads("".join(prop.itertext()))
            except ValueError:
                return None
            return plan if isinstance(plan, dict) and "total" in plan else None
    return None


def _read_plan(book: Path) -> dict:
    try:
        return json.loads(_plan_path(book).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        pass
    plan = _embedded_plan(book)  # the scratch copy was cleaned up
    if plan is None:
        raise WorkbookLiveError(f"No live plan for {book.name}; run `init` first.")
    return plan


def _plan_or_finished(book: Path, sheets: int) -> dict:
    """The workbook's plan, or a finished one for a workbook that has none."""
    try:
        return _read_plan(book)
    except WorkbookLiveError:
        return {"v": 1, "total": max(sheets, 1), "added": sheets, "state": "done"}


def _write(book: Path, workbook, plan: dict) -> None:
    buffer = io.BytesIO()
    workbook.save(buffer)
    _plan_path(book).write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    _atomic_write(book, _with_plan(buffer.getvalue(), plan))


def _open(book: Path):
    from openpyxl import load_workbook

    try:
        return load_workbook(book)
    except (OSError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise WorkbookLiveError(f"Cannot open {book}") from exc


def _sheet_count(book: Path) -> int:
    try:
        workbook = _open(book)
    except WorkbookLiveError:
        return 0
    plan = _embedded_plan(book) or {}
    if plan.get("added") == 0:
        return 0  # only the stand-in sheet
    return len(workbook.sheetnames)


def init(book: Path, sheets: int, *, force: bool = False) -> dict:
    """Create a workbook that expects ``sheets`` sheets.

    The count only sizes the preview's loading state: one skeleton per sheet
    not added yet. A workbook that already has sheets is left alone unless
    ``force`` starts it over.
    """
    from openpyxl import Workbook

    if not 0 < sheets <= _MAX_SHEETS:
        raise WorkbookLiveError(f"--sheets must be between 1 and {_MAX_SHEETS}.")
    existing = _sheet_count(book) if book.exists() else 0
    if existing and not force:
        raise WorkbookLiveError(
            f"{book.name} already has {existing} sheets. Fix a sheet with "
            "`add BOOK SHEET_FILE --replace N` or rebuild them all with `rebuild BOOK "
            "sheets/`; `init --force` empties the workbook and the preview starts over."
        )
    workbook = Workbook()
    workbook.active.title = _STAND_IN  # a workbook cannot have zero sheets
    plan = {"v": 1, "total": sheets, "added": 0, "state": "building"}
    # EvoFlux names the session running this command; only that session's
    # viewer shows the build live, and only while its turn runs.
    session = os.environ.get("EVOFLUX_SESSION", "").strip()
    if session:
        plan["session"] = session
    book.parent.mkdir(parents=True, exist_ok=True)
    _write(book, workbook, plan)
    return plan


def _forget_helpers(folder: Path) -> None:
    """Drop cached helper modules of the sheet folder (``theme.py`` …) so
    edits made since they were imported take effect."""
    local = {path.stem for path in folder.glob("*.py")}
    root = folder.resolve()
    for name, module in list(sys.modules.items()):
        location = getattr(module, "__file__", None)
        if name in local or (location and Path(location).resolve().parent == root):
            del sys.modules[name]
    importlib.invalidate_caches()


def _load_builder(script: Path):
    """Return the ``build`` function a sheet file defines."""
    import importlib.util

    if not script.is_file():
        raise WorkbookLiveError(f"No sheet file at {script}")
    folder = str(script.resolve().parent)
    if folder not in sys.path:
        sys.path.insert(0, folder)  # shared helpers beside the sheet files
    spec = importlib.util.spec_from_file_location(f"_evoflux_sheet_{script.stem}", script)
    if spec is None or spec.loader is None:
        raise WorkbookLiveError(f"Cannot load {script}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    build = getattr(module, "build", None)
    if not callable(build):
        raise WorkbookLiveError(f"{script.name} must define build(wb)")
    return build


def _build_one(workbook, script: Path):
    """Run ``script`` and return the one worksheet it added."""
    before = list(workbook.worksheets)
    _load_builder(script)(workbook)
    added = [sheet for sheet in workbook.worksheets if sheet not in before]
    if len(added) != 1 or len(workbook.worksheets) != len(before) + 1:
        raise WorkbookLiveError(
            f"{script.name} must add exactly one sheet (it added {len(added)})"
        )
    return added[0]


def add(book: Path, script: Path, *, replace: int | None = None) -> int:
    """Add the one sheet ``script`` builds (or rebuild sheet ``replace``).

    The workbook is saved only when the build succeeds, so a failing sheet
    file leaves the previous state on screen. Returns the sheet count.
    """
    workbook = _open(book)
    plan = _plan_or_finished(book, len(workbook.sheetnames))
    stand_in = plan.get("added") == 0 and workbook.sheetnames == [_STAND_IN]
    real = 0 if stand_in else len(workbook.sheetnames)
    if replace is not None and not 1 <= replace <= real:
        raise WorkbookLiveError(f"--replace {replace}: the workbook has {real} sheets")
    _forget_helpers(script.parent)
    old = workbook.worksheets[replace - 1] if replace is not None else None
    old_title = old.title if old is not None else None
    if old is not None:
        old.title = f"{old_title[:24]} (old)"  # free the name for the rebuild
    sheet = _build_one(workbook, script)
    if old is not None:
        workbook.move_sheet(sheet, offset=workbook.index(old) - workbook.index(sheet))
        workbook.remove(old)
    if stand_in:
        workbook.remove(workbook[_STAND_IN])
    workbook.active = workbook.index(sheet)
    count = len(workbook.sheetnames)
    plan["added"] = count
    _write(book, workbook, plan)
    return count


def rebuild(book: Path, folder: Path) -> int:
    """Rebuild every sheet from ``folder``'s numbered sheet files, in one save.

    For fixes that touch shared helpers or several sheets once the workbook
    is built: it keeps its plan (a finished workbook stays finished), so the
    preview redraws it instead of starting over. Returns the sheet count.
    """
    from openpyxl import Workbook

    scripts = sorted(path for path in folder.glob("*.py") if path.name[:1].isdigit())
    if not scripts:
        raise WorkbookLiveError(f"No numbered sheet files (01_inputs.py, …) in {folder}")
    plan = _plan_or_finished(book, len(scripts))
    workbook = Workbook()
    stand_in = workbook.active
    _forget_helpers(folder)
    for script in scripts:
        _build_one(workbook, script)
    workbook.remove(stand_in)
    workbook.active = 0
    plan["added"] = len(scripts)
    if plan.get("state") == "done":
        plan["total"] = len(scripts)
    _write(book, workbook, plan)
    return len(scripts)


def finish(book: Path) -> dict:
    """Mark the workbook complete; the plan stays so later fixes keep it so."""
    workbook = _open(book)
    plan = _plan_or_finished(book, len(workbook.sheetnames))
    plan["state"] = "done"
    _write(book, workbook, plan)
    return plan


def main() -> int:
    # Sheet files and helpers change between commands; never leave bytecode
    # beside them (stale caches, and __pycache__ folders among the user's files).
    sys.dont_write_bytecode = True
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n", 1)[0])
    commands = parser.add_subparsers(dest="command", required=True)
    start = commands.add_parser("init", help="create a workbook carrying the plan")
    start.add_argument("book", type=Path)
    start.add_argument("--sheets", type=int, required=True, help="how many sheets you will add")
    start.add_argument(
        "--force", action="store_true", help="empty a workbook that has sheets and start over"
    )
    one = commands.add_parser("add", help="add the one sheet a sheet file's build(wb) makes")
    one.add_argument("book", type=Path)
    one.add_argument("script", type=Path)
    one.add_argument("--replace", type=int, metavar="N", help="rebuild sheet N instead")
    end = commands.add_parser("finish", help="mark the workbook complete")
    end.add_argument("book", type=Path)
    redo = commands.add_parser("rebuild", help="rebuild every sheet from the sheet files")
    redo.add_argument("book", type=Path)
    redo.add_argument("folder", type=Path, help="folder with the numbered sheet files")
    args = parser.parse_args()
    try:
        if args.command == "init":
            plan = init(args.book, args.sheets, force=args.force)
            print(f"created {args.book} expecting {plan['total']} sheets")
        elif args.command == "add":
            count = add(args.book, args.script, replace=args.replace)
            plan = _read_plan(args.book)
            if args.replace is not None:
                print(f"rebuilt sheet {args.replace} of {args.book}")
            elif plan.get("state") == "done":
                print(f"sheet {count} saved; the workbook stays finished")
            elif count < plan["total"]:
                print(f"sheet {count} of {plan['total']} saved; write the next sheet file")
            else:
                print(f"sheet {count} of {plan['total']} saved; run `finish` when it is done")
        elif args.command == "rebuild":
            print(f"rebuilt {rebuild(args.book, args.folder)} sheets of {args.book}")
        else:
            finish(args.book)
            print(f"{args.book} is complete")
    except WorkbookLiveError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

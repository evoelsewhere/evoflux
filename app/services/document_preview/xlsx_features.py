"""Workbook features the openpyxl object model drops or the grid renderer skipped.

The grid renderer in :mod:`app.services.document_preview.service` draws cells
from openpyxl's object model. This module fills the gaps that made previews
look emptier than the workbook really is:

* conditional formatting (cell rules, colour scales, data bars, icon sets);
* drawing shapes and text boxes, which openpyxl discards on load;
* packages whose parts use ``mc:AlternateContent`` (for example Hancom Office
  styles), which openpyxl cannot parse.
"""

from __future__ import annotations

import html
import io
import math
import posixpath
import re
import zipfile
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree  # ty: ignore[unresolved-import] - compiled module, no stubs

_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_EMU_PER_PX = 914400 / 96
_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)

_THEME_COLOR_FALLBACK = {
    0: "#ffffff",
    1: "#000000",
    2: "#e7e6e6",
    3: "#44546a",
    4: "#4472c4",
    5: "#ed7d31",
    6: "#a5a5a5",
    7: "#ffc000",
    8: "#5b9bd5",
    9: "#70ad47",
}


# ---------------------------------------------------------------------------
# Markup-compatibility sanitising


def resolve_markup_compatibility(xml: bytes) -> bytes:
    """Replace every ``mc:AlternateContent`` with its ``mc:Fallback`` content."""
    if b"AlternateContent" not in xml:
        return xml
    root = etree.fromstring(xml, _XML_PARSER)
    for block in list(root.iter(f"{{{_MC_NS}}}AlternateContent")):
        parent = block.getparent()
        if parent is None:
            continue
        fallback = block.find(f"{{{_MC_NS}}}Fallback")
        index = parent.index(block)
        replacement = list(fallback) if fallback is not None else []
        tail = block.tail
        parent.remove(block)
        for offset, child in enumerate(replacement):
            parent.insert(index + offset, child)
        if tail and replacement:
            replacement[-1].tail = (replacement[-1].tail or "") + tail
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _name_unnamed_cell_styles(xml: bytes) -> bytes:
    if b"cellStyle" not in xml:
        return xml
    root = etree.fromstring(xml, _XML_PARSER)
    changed = False
    for index, style in enumerate(root.iter(f"{{{_SHEET_NS}}}cellStyle")):
        if not style.get("name"):
            style.set("name", f"Style {index + 1}")
            changed = True
    if not changed:
        return xml
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _drop_pivot_caches(xml: bytes) -> bytes:
    """Remove pivot cache references, which the grid preview never renders.

    openpyxl fails outright on some pivot cache definitions (for example ones
    with nested extension lists), taking the whole workbook down with them.
    """
    if b"pivotCaches" not in xml:
        return xml
    root = etree.fromstring(xml, _XML_PARSER)
    for caches in root.xpath("./*[local-name()='pivotCaches']"):
        root.remove(caches)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _drop_pivot_table_relationships(xml: bytes) -> bytes:
    """Detach worksheet pivot tables whose caches :func:`_drop_pivot_caches` removed."""
    if b"pivotTable" not in xml:
        return xml
    root = etree.fromstring(xml, _XML_PARSER)
    for relationship in list(root.iter(f"{{{_PACKAGE_REL_NS}}}Relationship")):
        if str(relationship.get("Type") or "").endswith("/pivotTable"):
            root.remove(relationship)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def sanitized_workbook_bytes(source: Path) -> bytes:
    """Return a copy of ``source`` whose XML parts openpyxl can load.

    The copy stays in memory: openpyxl reads embedded images lazily from the
    archive, so a temporary file would still be open (and undeletable on
    Windows) long after loading.
    """
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(source) as reader,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as writer,
    ):
        for member in reader.infolist():
            payload = reader.read(member.filename)
            try:
                if member.filename.endswith(".xml"):
                    payload = resolve_markup_compatibility(payload)
                    if member.filename == "xl/styles.xml":
                        payload = _name_unnamed_cell_styles(payload)
                    elif member.filename == "xl/workbook.xml":
                        payload = _drop_pivot_caches(payload)
                elif member.filename.startswith("xl/worksheets/_rels/"):
                    payload = _drop_pivot_table_relationships(payload)
            except etree.XMLSyntaxError:
                pass
            writer.writestr(member, payload)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Conditional formatting


@dataclass
class CellFormat:
    """Presentation produced by conditional formatting for one cell."""

    styles: dict[str, str] = field(default_factory=dict)
    icon: str = ""
    hide_value: bool = False

    def set(self, name: str, value: str) -> None:
        # Rules are applied in priority order; the first rule to set a
        # property wins, as in Excel.
        self.styles.setdefault(name, value)

    def css(self) -> list[str]:
        return [f"{name}:{value}" for name, value in self.styles.items()]


def numeric_value(value: Any) -> float | None:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)) and math.isfinite(value):
        return float(value)
    if isinstance(value, str):
        try:
            number = float(value.replace(",", ""))
        except ValueError:
            return None
        return number if math.isfinite(number) else None
    return None


def color_css(color: Any) -> str | None:
    """Resolve an openpyxl colour to CSS, including theme colours with tint."""
    if color is None:
        return None
    color_type = getattr(color, "type", None)
    rgb: str | None = None
    if color_type == "rgb" and isinstance(color.rgb, str):
        value = color.rgb[-6:]
        if re.fullmatch(r"[0-9A-Fa-f]{6}", value) and color.rgb != "00000000":
            rgb = f"#{value.lower()}"
    elif color_type == "theme":
        rgb = _THEME_COLOR_FALLBACK.get(int(color.theme or 0))
    elif color_type == "indexed":
        from openpyxl.styles.colors import COLOR_INDEX

        index = int(color.indexed or 0)
        if 0 <= index < len(COLOR_INDEX) and index not in (64, 65):
            rgb = f"#{COLOR_INDEX[index][-6:].lower()}"
    if rgb is None:
        return None
    tint = float(getattr(color, "tint", 0) or 0)
    return _apply_tint(rgb, tint) if tint else rgb


def _apply_tint(rgb: str, tint: float) -> str:
    channels = [int(rgb[index : index + 2], 16) for index in (1, 3, 5)]
    if tint < 0:
        channels = [round(channel * (1 + tint)) for channel in channels]
    else:
        channels = [round(channel + (255 - channel) * tint) for channel in channels]
    return "#" + "".join(f"{max(0, min(255, c)):02x}" for c in channels)


def _mix(start: str, end: str, ratio: float) -> str:
    ratio = max(0.0, min(1.0, ratio))
    first = [int(start[index : index + 2], 16) for index in (1, 3, 5)]
    second = [int(end[index : index + 2], 16) for index in (1, 3, 5)]
    return "#" + "".join(
        f"{round(a + (b - a) * ratio):02x}" for a, b in zip(first, second, strict=True)
    )


def _dxf_styles(dxf: Any) -> dict[str, str]:
    styles: dict[str, str] = {}
    if dxf is None:
        return styles
    fill = getattr(dxf, "fill", None)
    if fill is not None:
        # Differential fills keep the visible colour in bgColor for solid
        # patterns, but many writers only populate fgColor.
        color = color_css(getattr(fill, "bgColor", None)) or color_css(
            getattr(fill, "fgColor", None)
        )
        if color:
            styles["background"] = color
    font = getattr(dxf, "font", None)
    if font is not None:
        color = color_css(getattr(font, "color", None))
        if color:
            styles["color"] = color
        if getattr(font, "b", None):
            styles["font-weight"] = "700"
        if getattr(font, "i", None):
            styles["font-style"] = "italic"
        if getattr(font, "strike", None):
            styles["text-decoration"] = "line-through"
        elif getattr(font, "u", None):
            styles["text-decoration"] = "underline"
    return styles


def _formula_number(formula: list[str] | None, index: int = 0) -> float | None:
    if not formula or len(formula) <= index:
        return None
    return numeric_value(str(formula[index]).strip())


def _formula_text(formula: list[str] | None, text: str | None) -> str:
    if text:
        return text
    if formula:
        match = re.search(r'"([^"]*)"', formula[0])
        if match:
            return match.group(1)
    return ""


def _cell_is(operator: str | None, value: Any, formula: list[str] | None) -> bool:
    number = numeric_value(value)
    first = _formula_number(formula)
    second = _formula_number(formula, 1)
    if number is None or first is None:
        if operator in {"equal", "notEqual"} and formula:
            expected = str(formula[0]).strip().strip('"')
            matched = str(value if value is not None else "") == expected
            return matched if operator == "equal" else not matched
        return False
    if operator == "between" and second is not None:
        return min(first, second) <= number <= max(first, second)
    if operator == "notBetween" and second is not None:
        return not min(first, second) <= number <= max(first, second)
    return {
        "equal": number == first,
        "notEqual": number != first,
        "greaterThan": number > first,
        "greaterThanOrEqual": number >= first,
        "lessThan": number < first,
        "lessThanOrEqual": number <= first,
    }.get(operator or "", False)


def _threshold(cfvo: Any, numbers: list[float], default: float) -> float:
    if not numbers:
        return default
    kind = getattr(cfvo, "type", None)
    raw = numeric_value(getattr(cfvo, "val", None))
    if kind == "min":
        return min(numbers)
    if kind == "max":
        return max(numbers)
    if kind == "percent" and raw is not None:
        low, high = min(numbers), max(numbers)
        return low + (high - low) * raw / 100
    if kind == "percentile" and raw is not None:
        ordered = sorted(numbers)
        position = (len(ordered) - 1) * raw / 100
        lower = math.floor(position)
        upper = math.ceil(position)
        return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)
    if raw is not None:
        return raw
    return default


_ICON_SETS = {
    "3Arrows": ("#e74c3c:▼", "#f1c40f:►", "#27ae60:▲"),
    "3ArrowsGray": ("#7f8c8d:▼", "#7f8c8d:►", "#7f8c8d:▲"),
    "3TrafficLights1": ("#e74c3c:●", "#f1c40f:●", "#27ae60:●"),
    "3TrafficLights2": ("#e74c3c:●", "#f1c40f:●", "#27ae60:●"),
    "3Signs": ("#e74c3c:◆", "#f1c40f:▲", "#27ae60:●"),
    "3Symbols": ("#e74c3c:✖", "#f1c40f:!", "#27ae60:✔"),
    "3Symbols2": ("#e74c3c:✖", "#f1c40f:!", "#27ae60:✔"),
    "3Flags": ("#e74c3c:⚑", "#f1c40f:⚑", "#27ae60:⚑"),
    "4Arrows": ("#e74c3c:▼", "#e67e22:◢", "#f1c40f:◥", "#27ae60:▲"),
    "4TrafficLights": ("#2c3e50:●", "#e74c3c:●", "#f1c40f:●", "#27ae60:●"),
    "5Arrows": ("#e74c3c:▼", "#e67e22:◢", "#f1c40f:►", "#9acd32:◥", "#27ae60:▲"),
    "5Quarters": ("#bdc3c7:○", "#7f8c8d:◔", "#7f8c8d:◑", "#7f8c8d:◕", "#2c3e50:●"),
    "5Rating": ("#bdc3c7:▁", "#3498db:▂", "#3498db:▃", "#3498db:▅", "#3498db:▇"),
}


def conditional_formats(
    sheet: Any, value_of: Callable[[str], Any]
) -> dict[str, CellFormat]:
    """Evaluate the sheet's conditional formatting into per-cell presentation."""
    from openpyxl.utils import get_column_letter

    entries: list[tuple[int, list[str], Any]] = []
    for formatting in sheet.conditional_formatting:
        coordinates: list[str] = []
        for cell_range in formatting.sqref.ranges:
            rows = range(cell_range.min_row, min(cell_range.max_row, 5000) + 1)
            columns = range(cell_range.min_col, min(cell_range.max_col, 200) + 1)
            coordinates.extend(
                f"{get_column_letter(column)}{row}"
                for row in rows
                for column in columns
            )
        for rule in formatting.rules:
            entries.append((int(rule.priority or 0), coordinates, rule))
    entries.sort(key=lambda entry: entry[0])

    result: dict[str, CellFormat] = {}
    for _priority, coordinates, rule in entries:
        values = {coordinate: value_of(coordinate) for coordinate in coordinates}
        numbers = [
            number
            for number in (numeric_value(value) for value in values.values())
            if number is not None
        ]
        kind = rule.type
        if kind == "colorScale" and rule.colorScale is not None:
            _apply_color_scale(rule.colorScale, values, numbers, result)
            continue
        if kind == "dataBar" and rule.dataBar is not None:
            _apply_data_bar(rule.dataBar, values, numbers, result)
            continue
        if kind == "iconSet" and rule.iconSet is not None:
            _apply_icon_set(rule.iconSet, values, numbers, result)
            continue
        matched = _matching_cells(rule, values, numbers)
        styles = _dxf_styles(rule.dxf)
        if not styles:
            continue
        for coordinate in matched:
            target = result.setdefault(coordinate, CellFormat())
            for name, value in styles.items():
                target.set(name, value)
    return result


def _matching_cells(
    rule: Any, values: dict[str, Any], numbers: list[float]
) -> list[str]:
    kind = rule.type
    formula = list(rule.formula or [])
    if kind == "cellIs":
        return [c for c, v in values.items() if _cell_is(rule.operator, v, formula)]
    if kind in {"containsText", "notContainsText", "beginsWith", "endsWith"}:
        needle = _formula_text(formula, rule.text).casefold()
        if not needle:
            return []
        checks = {
            "containsText": lambda text: needle in text,
            "notContainsText": lambda text: needle not in text,
            "beginsWith": lambda text: text.startswith(needle),
            "endsWith": lambda text: text.endswith(needle),
        }
        check = checks[kind]
        return [
            coordinate
            for coordinate, value in values.items()
            if value is not None and check(str(value).casefold())
        ]
    if kind in {"containsBlanks", "notContainsBlanks"}:
        blank = kind == "containsBlanks"
        return [
            coordinate
            for coordinate, value in values.items()
            if (value is None or str(value).strip() == "") == blank
        ]
    if kind == "top10" and numbers:
        rank = int(rule.rank or 10)
        count = max(1, math.floor(len(numbers) * rank / 100)) if rule.percent else rank
        ordered = sorted(numbers, reverse=not rule.bottom)
        cutoff = ordered[min(count, len(ordered)) - 1]
        return [
            coordinate
            for coordinate, value in values.items()
            if (number := numeric_value(value)) is not None
            and (number <= cutoff if rule.bottom else number >= cutoff)
        ]
    if kind == "aboveAverage" and numbers:
        mean = sum(numbers) / len(numbers)
        above = rule.aboveAverage is not False
        inclusive = bool(rule.equalAverage)
        matched: list[str] = []
        for coordinate, value in values.items():
            number = numeric_value(value)
            if number is None:
                continue
            if (inclusive and number == mean) or (
                number > mean if above else number < mean
            ):
                matched.append(coordinate)
        return matched
    if kind in {"duplicateValues", "uniqueValues"}:
        counts = Counter(
            str(value) for value in values.values() if value not in (None, "")
        )
        duplicate = kind == "duplicateValues"
        return [
            coordinate
            for coordinate, value in values.items()
            if value not in (None, "") and (counts[str(value)] > 1) == duplicate
        ]
    return []


def _apply_color_scale(
    scale: Any,
    values: dict[str, Any],
    numbers: list[float],
    result: dict[str, CellFormat],
) -> None:
    colors = [color_css(color) or "#ffffff" for color in scale.color]
    if len(colors) < 2 or not numbers:
        return
    stops = [
        _threshold(cfvo, numbers, numbers[0]) for cfvo in scale.cfvo[: len(colors)]
    ]
    for coordinate, value in values.items():
        number = numeric_value(value)
        if number is None:
            continue
        if len(stops) == 3 and len(colors) >= 3:
            if number <= stops[1]:
                span = stops[1] - stops[0]
                color = _mix(
                    colors[0], colors[1], (number - stops[0]) / span if span else 0
                )
            else:
                span = stops[2] - stops[1]
                color = _mix(
                    colors[1], colors[2], (number - stops[1]) / span if span else 1
                )
        else:
            span = stops[-1] - stops[0]
            color = _mix(
                colors[0], colors[-1], (number - stops[0]) / span if span else 1
            )
        result.setdefault(coordinate, CellFormat()).set("background", color)


def _apply_data_bar(
    bar: Any,
    values: dict[str, Any],
    numbers: list[float],
    result: dict[str, CellFormat],
) -> None:
    if not numbers:
        return
    low = _threshold(bar.cfvo[0], numbers, min(numbers)) if bar.cfvo else min(numbers)
    high = (
        _threshold(bar.cfvo[1], numbers, max(numbers))
        if len(bar.cfvo) > 1
        else max(numbers)
    )
    color = color_css(bar.color) or "#638ec6"
    minimum = float(bar.minLength if bar.minLength is not None else 10)
    maximum = float(bar.maxLength if bar.maxLength is not None else 90)
    for coordinate, value in values.items():
        number = numeric_value(value)
        if number is None:
            continue
        ratio = (number - low) / (high - low) if high > low else 1
        length = minimum + (maximum - minimum) * max(0.0, min(1.0, ratio))
        target = result.setdefault(coordinate, CellFormat())
        target.set(
            "background",
            f"linear-gradient(90deg,{color} 0,{color}33 {length:.1f}%,"
            f"transparent {length:.1f}%)",
        )
        if bar.showValue is False:
            target.hide_value = True


def _apply_icon_set(
    icon_set: Any,
    values: dict[str, Any],
    numbers: list[float],
    result: dict[str, CellFormat],
) -> None:
    icons = _ICON_SETS.get(str(icon_set.iconSet or "3TrafficLights1"))
    if not icons or not numbers:
        return
    if icon_set.reverse:
        icons = tuple(reversed(icons))
    thresholds = [
        _threshold(cfvo, numbers, min(numbers)) for cfvo in icon_set.cfvo[: len(icons)]
    ]
    for coordinate, value in values.items():
        number = numeric_value(value)
        if number is None:
            continue
        level = 0
        for index, threshold in enumerate(thresholds):
            if index and number >= threshold:
                level = index
        color, glyph = icons[min(level, len(icons) - 1)].split(":", 1)
        target = result.setdefault(coordinate, CellFormat())
        if not target.icon:
            target.icon = (
                f'<span class="cf-icon" style="color:{color}" aria-hidden="true">'
                f"{glyph}</span>"
            )
        if icon_set.showValue is False:
            target.hide_value = True


# ---------------------------------------------------------------------------
# Drawing shapes and text boxes


def _relationship_targets(archive: zipfile.ZipFile, part: str) -> dict[str, str]:
    directory, name = posixpath.split(part)
    rels_path = posixpath.join(directory, "_rels", f"{name}.rels")
    try:
        root = etree.fromstring(archive.read(rels_path), _XML_PARSER)
    except (KeyError, etree.XMLSyntaxError):
        return {}
    targets: dict[str, str] = {}
    for relationship in root.iter(f"{{{_PACKAGE_REL_NS}}}Relationship"):
        if relationship.get("TargetMode") == "External":
            continue
        target = relationship.get("Target") or ""
        resolved = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join(directory, target))
        )
        targets[relationship.get("Id") or ""] = resolved
    return targets


def visible_sheet_names(source: Path) -> list[str]:
    """Return workbook sheet names in tab order, skipping hidden sheets."""
    try:
        with zipfile.ZipFile(source) as archive:
            workbook = etree.fromstring(archive.read("xl/workbook.xml"), _XML_PARSER)
    except (KeyError, OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        return []
    return [
        str(sheet.get("name") or "")
        for sheet in workbook.iter(f"{{{_SHEET_NS}}}sheet")
        if (sheet.get("state") or "visible") == "visible"
    ]


def worksheet_drawing_parts(source: Path) -> dict[str, str]:
    """Map worksheet titles to the drawing part each one references."""
    result: dict[str, str] = {}
    try:
        with zipfile.ZipFile(source) as archive:
            workbook = etree.fromstring(archive.read("xl/workbook.xml"), _XML_PARSER)
            sheet_targets = _relationship_targets(archive, "xl/workbook.xml")
            for sheet in workbook.iter(f"{{{_SHEET_NS}}}sheet"):
                part = sheet_targets.get(sheet.get(f"{{{_REL_NS}}}id") or "")
                if not part or part not in archive.namelist():
                    continue
                sheet_root = etree.fromstring(archive.read(part), _XML_PARSER)
                drawing = sheet_root.find(f"{{{_SHEET_NS}}}drawing")
                if drawing is None:
                    continue
                drawing_part = _relationship_targets(archive, part).get(
                    drawing.get(f"{{{_REL_NS}}}id") or ""
                )
                if drawing_part:
                    result[str(sheet.get("name"))] = drawing_part
    except (KeyError, OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        return {}
    return result


def _local(element: Any) -> str:
    return str(element.tag).rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""


def _child(element: Any | None, name: str) -> Any | None:
    if element is None:
        return None
    for child in element:
        if _local(child) == name:
            return child
    return None


def _drawing_color(element: Any | None) -> str | None:
    if element is None:
        return None
    for child in element:
        name = _local(child)
        if name == "srgbClr":
            value = child.get("val") or ""
            if re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                return f"#{value.lower()}"
        if name == "schemeClr":
            index = {
                "lt1": 0,
                "bg1": 0,
                "dk1": 1,
                "tx1": 1,
                "lt2": 2,
                "bg2": 2,
                "dk2": 3,
                "tx2": 3,
                "accent1": 4,
                "accent2": 5,
                "accent3": 6,
                "accent4": 7,
                "accent5": 8,
                "accent6": 9,
            }.get(child.get("val") or "")
            if index is not None:
                return _THEME_COLOR_FALLBACK[index]
        if name == "sysClr":
            value = child.get("lastClr") or ""
            if re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                return f"#{value.lower()}"
    return None


def _marker_px(
    marker: Any,
    column_left: Callable[[int], float],
    row_top: Callable[[int], float],
) -> tuple[float, float]:
    def number(name: str) -> int:
        node = _child(marker, name)
        try:
            return int((node.text or "0") if node is not None else "0")
        except ValueError:
            return 0

    return (
        column_left(number("col")) + number("colOff") / _EMU_PER_PX,
        row_top(number("row")) + number("rowOff") / _EMU_PER_PX,
    )


def _anchor_box(
    anchor: Any,
    column_left: Callable[[int], float],
    row_top: Callable[[int], float],
) -> tuple[float, float, float, float] | None:
    kind = _local(anchor)
    start = _child(anchor, "from")
    if kind == "absoluteAnchor":
        position, extent = _child(anchor, "pos"), _child(anchor, "ext")
        if position is None or extent is None:
            return None
        left = 44 + int(position.get("x", "0")) / _EMU_PER_PX
        top = 25 + int(position.get("y", "0")) / _EMU_PER_PX
        return (
            left,
            top,
            int(extent.get("cx", "0")) / _EMU_PER_PX,
            int(extent.get("cy", "0")) / _EMU_PER_PX,
        )
    if start is None:
        return None
    left, top = _marker_px(start, column_left, row_top)
    if kind == "twoCellAnchor":
        end = _child(anchor, "to")
        if end is None:
            return None
        right, bottom = _marker_px(end, column_left, row_top)
        return left, top, max(right - left, 1), max(bottom - top, 1)
    extent = _child(anchor, "ext")
    if extent is None:
        return None
    return (
        left,
        top,
        int(extent.get("cx", "0")) / _EMU_PER_PX,
        int(extent.get("cy", "0")) / _EMU_PER_PX,
    )


_ROUND_GEOMETRIES = frozenset({"ellipse", "flowChartConnector", "donut"})
_ROUNDED_GEOMETRIES = frozenset({"roundRect", "flowChartAlternateProcess"})


def _shape_html(shape: Any, box: tuple[float, float, float, float]) -> str:
    left, top, width, height = box
    properties = _child(shape, "spPr")
    styles = [
        f"left:{left:.2f}px",
        f"top:{top:.2f}px",
        f"width:{width:.2f}px",
        f"height:{height:.2f}px",
    ]
    if properties is not None:
        fill = _child(properties, "solidFill")
        if fill is not None:
            color = _drawing_color(fill)
            if color:
                styles.append(f"background:{color}")
        line = _child(properties, "ln")
        if line is not None and _child(line, "noFill") is None:
            color = _drawing_color(_child(line, "solidFill")) or "#4472c4"
            width_emu = int(line.get("w", "9525") or "9525")
            styles.append(
                f"border:{max(width_emu / _EMU_PER_PX, 1):.1f}px solid {color}"
            )
        geometry = _child(properties, "prstGeom")
        preset = geometry.get("prst") if geometry is not None else ""
        if preset in _ROUND_GEOMETRIES:
            styles.append("border-radius:50%")
        elif preset in _ROUNDED_GEOMETRIES:
            styles.append("border-radius:12px")
        transform = _child(properties, "xfrm")
        rotation = int(transform.get("rot", "0") or "0") if transform is not None else 0
        if rotation:
            styles.append(f"transform:rotate({rotation / 60000:.2f}deg)")
    paragraphs: list[str] = []
    body = _child(shape, "txBody")
    if body is not None:
        for paragraph in body:
            if _local(paragraph) != "p":
                continue
            runs: list[str] = []
            align = ""
            for node in paragraph:
                name = _local(node)
                if name == "pPr":
                    align = {"ctr": "center", "r": "right", "just": "justify"}.get(
                        node.get("algn") or "", ""
                    )
                elif name in {"r", "fld"}:
                    text = _child(node, "t")
                    run_styles: list[str] = []
                    run_properties = _child(node, "rPr")
                    if run_properties is not None:
                        if run_properties.get("b") == "1":
                            run_styles.append("font-weight:700")
                        if run_properties.get("i") == "1":
                            run_styles.append("font-style:italic")
                        size = run_properties.get("sz")
                        if size and size.isdigit():
                            run_styles.append(f"font-size:{int(size) / 100:.1f}pt")
                        color = _drawing_color(_child(run_properties, "solidFill"))
                        if color:
                            run_styles.append(f"color:{color}")
                    style = f' style="{";".join(run_styles)}"' if run_styles else ""
                    runs.append(
                        f"<span{style}>{html.escape(text.text or '' if text is not None else '')}</span>"
                    )
                elif name == "br":
                    runs.append("<br>")
            style = f' style="text-align:{align}"' if align else ""
            paragraphs.append(f"<p{style}>{''.join(runs) or '&nbsp;'}</p>")
    properties_node = _child(shape, "nvSpPr")
    name_node = _child(properties_node, "cNvPr")
    name = html.escape(
        str(name_node.get("name") if name_node is not None else ""), True
    )
    shape_id = html.escape(
        str(name_node.get("id") if name_node is not None else ""), True
    )
    return (
        f'<div class="workbook-shape" data-shape-id="{shape_id}" '
        f'data-shape-name="{name}" data-qa-label="{name or "Shape"}" '
        f'style="{";".join(styles)}">{"".join(paragraphs)}</div>'
    )


def drawing_shapes_html(
    source: Path,
    drawing_part: str,
    column_left: Callable[[int], float],
    row_top: Callable[[int], float],
) -> tuple[list[str], float, float]:
    """Render the shapes/text boxes of a worksheet drawing.

    Returns the HTML fragments plus the right and bottom edges they occupy,
    so the caller can grow its stage to fit them.
    """
    try:
        with zipfile.ZipFile(source) as archive:
            root = etree.fromstring(archive.read(drawing_part), _XML_PARSER)
    except (KeyError, OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        return [], 0, 0
    fragments: list[str] = []
    right = bottom = 0.0
    for anchor in root:
        if _local(anchor) not in {"twoCellAnchor", "oneCellAnchor", "absoluteAnchor"}:
            continue
        box = _anchor_box(anchor, column_left, row_top)
        if box is None:
            continue
        for shape in anchor.iter():
            if _local(shape) != "sp":
                continue
            fragments.append(_shape_html(shape, box))
            right = max(right, box[0] + box[2])
            bottom = max(bottom, box[1] + box[3])
            break
    return fragments, right, bottom

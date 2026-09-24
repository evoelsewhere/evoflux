"""Routes for the version history of workspace Office documents."""

from __future__ import annotations

import uuid

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

pytestmark = pytest.mark.usefixtures("setup_db")


@pytest.fixture
def client():
    from app.api.app import create_app
    from app.services.team_manager import set_team

    app = create_app()
    set_team(None)
    yield TestClient(app)
    set_team(None)


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    from app.api.routes.team import files as team_files
    from app.services import document_versions

    root = tmp_path / "ws"
    root.mkdir()
    monkeypatch.setattr(team_files, "workspace_dir", lambda sid: root)
    monkeypatch.setattr(
        document_versions.settings, "EVOFLUX_STATE_DIR", str(tmp_path / "state")
    )

    async def no_watch(*_args):
        return None

    monkeypatch.setattr(document_versions, "ensure_watching", no_watch)
    return root


def _deck(path, text: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = text
    prs.save(path)


def test_history_checkpoint_and_undo(client, workspace):
    session_id = str(uuid.uuid7())
    url = f"/api/team/{session_id}/document-versions"
    _deck(workspace / "deck.pptx", "one")

    history = client.get(url, params={"path": "deck.pptx"}).json()
    assert len(history["versions"]) == 1 and not history["can_undo"]

    client.post(
        url, json={"path": "deck.pptx", "action": "checkpoint", "label": "Blue"}
    )
    _deck(workspace / "deck.pptx", "two")
    history = client.get(url, params={"path": "deck.pptx"}).json()
    assert history["versions"][-1]["label"] == "Blue"
    assert history["can_undo"]

    undone = client.post(url, json={"path": "deck.pptx", "action": "undo"})
    assert undone.status_code == 200 and undone.json()["can_redo"]
    nothing = client.post(url, json={"path": "deck.pptx", "action": "undo"})
    assert nothing.status_code == 409


def test_rejects_other_files_and_escapes(client, workspace):
    session_id = str(uuid.uuid7())
    url = f"/api/team/{session_id}/document-versions"
    (workspace / "notes.txt").write_text("hi")

    assert client.get(url, params={"path": "notes.txt"}).status_code == 415
    assert client.get(url, params={"path": "../deck.pptx"}).status_code == 400
    assert client.get(url, params={"path": "missing.pptx"}).status_code == 404

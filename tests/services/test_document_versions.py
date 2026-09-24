from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.services import document_versions as versions
from app.services.document_versions import DocumentVersionError

SESSION = "00000000-0000-0000-0000-000000000001"
TARGETS = (
    Path(__file__).resolve().parents[2]
    / "app/agent/builtin_skills/office-annotation-edit/scripts/targets.py"
)


@pytest.fixture(autouse=True)
def state_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(versions.settings, "EVOFLUX_STATE_DIR", str(tmp_path / "state"))
    versions._pending_labels.clear()


def _deck(path: Path, text: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = text
    prs.save(path)


def test_records_changes_and_steps_back_and_forward(tmp_path):
    root = tmp_path / "ws"
    root.mkdir()
    deck = root / "deck.pptx"
    _deck(deck, "one")

    history = versions.get_history(SESSION, root, "deck.pptx")
    assert [v.source for v in history.versions] == ["baseline"]
    assert not history.can_undo

    versions.checkpoint(SESSION, root, "deck.pptx", "Make it blue")
    _deck(deck, "two")
    history = versions.record_version(SESSION, root, "deck.pptx")
    assert len(history.versions) == 2
    assert history.versions[-1].label == "Make it blue"
    assert history.can_undo and not history.can_redo
    two = deck.read_bytes()

    history = versions.step(SESSION, root, "deck.pptx", offset=-1)
    assert history.can_redo
    assert "one" in Presentation(str(deck)).slides[0].shapes[0].text_frame.text
    # The restored file is already the head: the watcher records nothing new.
    assert len(versions.record_version(SESSION, root, "deck.pptx").versions) == 2

    history = versions.step(SESSION, root, "deck.pptx", offset=1)
    assert deck.read_bytes() == two
    with pytest.raises(DocumentVersionError):
        versions.step(SESSION, root, "deck.pptx", offset=1)


def test_a_new_save_after_undo_replaces_the_redo_branch(tmp_path):
    root = tmp_path / "ws"
    root.mkdir()
    deck = root / "deck.pptx"
    for text in ("one", "two"):
        _deck(deck, text)
        versions.record_version(SESSION, root, "deck.pptx")
    versions.step(SESSION, root, "deck.pptx", offset=-1)

    _deck(deck, "three")
    history = versions.record_version(SESSION, root, "deck.pptx")

    assert len(history.versions) == 2
    assert not history.can_redo
    blobs = list((versions._history_dir(SESSION, "deck.pptx") / "blobs").iterdir())
    assert len(blobs) == 2  # the dropped version's bytes are gone too


def test_ignores_half_written_files_and_live_decks(tmp_path, monkeypatch):
    root = tmp_path / "ws"
    root.mkdir()
    deck = root / "deck.pptx"
    deck.write_bytes(b"PK\x03\x04 not finished")
    assert versions.get_history(SESSION, root, "deck.pptx").versions == []

    _deck(deck, "one")
    monkeypatch.setattr(versions, "_being_built", lambda _path: True)
    assert versions.get_history(SESSION, root, "deck.pptx").versions == []
    assert not versions.is_versioned("~$deck.pptx")


def _chart_deck(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(8), Inches(0.8)
    ).text = "Title"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", (1.0, 2.0))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4), data
    )
    prs.save(path)


def test_targets_resolves_shapes_by_id_and_by_area(tmp_path, capsys):
    spec = importlib.util.spec_from_file_location("annotation_targets", TARGETS)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules["annotation_targets"] = module
    spec.loader.exec_module(module)
    deck = tmp_path / "deck.pptx"
    _chart_deck(deck)
    chart_id = Presentation(str(deck)).slides[0].shapes[1].shape_id

    assert module.main([str(deck), "--slide", "1", "--shape", str(chart_id)]) == 0
    by_id = json.loads(capsys.readouterr().out)
    assert by_id["targets"][0]["kind"] == "chart"
    assert by_id["targets"][0]["chart"]["series"][0]["name"] == "Revenue"

    assert module.main([str(deck), "--slide", "1", "--area", "5,25,90,70"]) == 0
    by_area = json.loads(capsys.readouterr().out)
    assert [item["kind"] for item in by_area["in_area"]] == ["chart"]
    sys.modules.pop("annotation_targets", None)

from __future__ import annotations

import asyncio
import io
import json
import subprocess
import tarfile
from pathlib import Path

import pytest

from app.services.office_runtime import convert, installer, manifest

PLATFORM = "win32-x64"
EXECUTABLE = "soffice/program/soffice.exe"


def _archive(
    path: Path, members: dict[str, bytes], links: dict[str, str] | None = None
) -> None:
    with tarfile.open(path, "w:gz") as archive:
        for name, payload in members.items():
            info = tarfile.TarInfo(name)
            info.size = len(payload)
            archive.addfile(info, io.BytesIO(payload))
        for name, target in (links or {}).items():
            info = tarfile.TarInfo(name)
            info.type = tarfile.SYMTYPE
            info.linkname = target
            archive.addfile(info)


def _record(archive: Path) -> dict:
    import hashlib

    return {
        "version": "26.8.0",
        "url": archive.resolve().as_uri(),
        "sha256": hashlib.sha256(archive.read_bytes()).hexdigest(),
        "size": archive.stat().st_size,
        "executable": EXECUTABLE,
    }


@pytest.fixture
def runtime_env(monkeypatch, tmp_path):
    monkeypatch.setattr(manifest, "platform_key", lambda: PLATFORM)
    monkeypatch.setattr(installer.settings, "EVOFLUX_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setattr(installer, "_job", None)
    monkeypatch.setattr(installer, "_task", None)

    def configure(record: dict) -> None:
        override = tmp_path / "manifest.json"
        override.write_text(json.dumps({"assets": {PLATFORM: record}}))
        monkeypatch.setenv(manifest.MANIFEST_OVERRIDE_ENV, str(override))

    return configure


def test_parse_asset_rejects_malformed_records():
    record = {
        "version": "26.8.0",
        "url": "https://example.com/libreoffice.tar.gz",
        "sha256": "a" * 64,
        "size": 1024,
        "executable": EXECUTABLE,
    }

    asset = manifest.parse_asset(PLATFORM, record)
    assert asset.size == 1024

    with pytest.raises(manifest.RuntimeManifestError, match="sha256"):
        manifest.parse_asset(PLATFORM, {**record, "sha256": "not-a-hash"})
    with pytest.raises(manifest.RuntimeManifestError, match="size"):
        manifest.parse_asset(PLATFORM, {**record, "size": 0})
    with pytest.raises(manifest.RuntimeManifestError, match="HTTPS"):
        manifest.parse_asset(PLATFORM, {**record, "url": "http://example.com/x"})
    with pytest.raises(manifest.RuntimeManifestError, match="soffice/"):
        manifest.parse_asset(PLATFORM, {**record, "executable": "../evil.exe"})


@pytest.mark.parametrize(
    ("members", "links", "message"),
    [
        ({"../evil": b"x"}, None, "escapes"),
        ({"other/file": b"x"}, None, "escapes"),
        ({"soffice/a": b"x"}, {"soffice/link": "../../etc/passwd"}, "link escapes"),
    ],
)
def test_extraction_refuses_unsafe_archives(tmp_path, members, links, message):
    archive = tmp_path / "bad.tar.gz"
    _archive(archive, members, links)

    with pytest.raises(installer.RuntimeInstallError, match=message):
        installer._extract(archive, tmp_path / "out")
    assert not (tmp_path / "out" / "evil").exists()


def test_install_verifies_extracts_and_activates(runtime_env, tmp_path):
    archive = tmp_path / "runtime.tar.gz"
    _archive(archive, {EXECUTABLE: b"MZ", "soffice/share/fonts/a.ttf": b"font"})
    runtime_env(_record(archive))

    async def run() -> None:
        installer.start_runtime_install()
        assert installer._task is not None
        await installer._task

    asyncio.run(run())

    status = installer.runtime_status()
    assert status.job is None
    assert status.installed_version == "26.8.0"
    runtime = installer.installed_runtime()
    assert runtime is not None
    assert runtime.executable.read_bytes() == b"MZ"
    assert not list(installer.runtime_home().glob(".staging-*"))


def test_install_rejects_archive_that_does_not_match_pin(runtime_env, tmp_path):
    archive = tmp_path / "runtime.tar.gz"
    _archive(archive, {EXECUTABLE: b"MZ"})
    record = _record(archive)
    _archive(archive, {EXECUTABLE: b"MZ-tampered-and-longer"})
    runtime_env(record)

    async def run() -> None:
        installer.start_runtime_install()
        assert installer._task is not None
        await installer._task

    asyncio.run(run())

    job = installer.runtime_status().job
    assert job is not None and job.phase == "failed"
    assert installer.installed_runtime() is None


def test_start_install_without_published_asset_is_refused(runtime_env, monkeypatch):
    monkeypatch.delenv(manifest.MANIFEST_OVERRIDE_ENV, raising=False)
    monkeypatch.setattr(manifest, "PINNED_ASSETS", {})
    with pytest.raises(installer.RuntimeInstallError, match="No verified"):
        installer.start_runtime_install()


def test_conversion_retries_the_silent_first_run_of_a_fresh_profile(
    monkeypatch, tmp_path
):
    executable = tmp_path / "soffice.exe"
    executable.write_bytes(b"")
    monkeypatch.setattr(
        convert,
        "installed_runtime",
        lambda: installer.InstalledRuntime("26.8.0", executable, tmp_path),
    )
    monkeypatch.setattr(convert, "runtime_home", lambda: tmp_path / "home")
    source = tmp_path / "deck.pptx"
    source.write_bytes(b"pptx")
    calls: list[list[str]] = []

    def fake_run(command, environment, timeout):
        calls.append(command)
        if len(calls) == 2:
            outdir = Path(command[command.index("--outdir") + 1])
            (outdir / "document.pdf").write_bytes(b"%PDF-1.4")
        return 0, b""

    monkeypatch.setattr(convert, "_run", fake_run)

    pdf = convert.convert_to_pdf(source, tmp_path / "work")

    assert pdf.read_bytes() == b"%PDF-1.4"
    assert len(calls) == 2
    assert "pdf:impress_pdf_Export" in calls[0]
    profile = (
        tmp_path / "home" / "profile" / "user" / "registrymodifications.xcu"
    ).read_text()
    assert "DisableMacrosExecution" in profile
    assert "BlockUntrustedRefererLinks" in profile


def test_conversion_reports_failure_without_output(monkeypatch, tmp_path):
    executable = tmp_path / "soffice.exe"
    executable.write_bytes(b"")
    monkeypatch.setattr(
        convert,
        "installed_runtime",
        lambda: installer.InstalledRuntime("26.8.0", executable, tmp_path),
    )
    monkeypatch.setattr(convert, "runtime_home", lambda: tmp_path / "home")
    monkeypatch.setattr(convert, "_run", lambda *args: (1, b"boom"))
    source = tmp_path / "book.xlsx"
    source.write_bytes(b"xlsx")

    with pytest.raises(convert.ConversionError):
        convert.convert_to_pdf(source, tmp_path / "work")


def test_conversion_kills_the_process_tree_on_timeout(monkeypatch, tmp_path):
    killed: list[int] = []

    class SlowProcess:
        pid = 4242
        returncode = None

        def communicate(self, timeout):
            raise subprocess.TimeoutExpired("soffice", timeout)

    monkeypatch.setattr(convert.subprocess, "Popen", lambda *a, **k: SlowProcess())
    monkeypatch.setattr(
        convert, "_kill_tree", lambda process: killed.append(process.pid)
    )

    with pytest.raises(convert.ConversionError, match="did not finish"):
        convert._run(["soffice"], {}, 1)
    assert killed == [4242]


def _text_pdf(path: Path, lines: list[tuple[str, int]]) -> None:
    """Write a one-page PDF (612x792pt) with Helvetica text at given y offsets."""
    content = "".join(
        f"BT /F1 24 Tf 72 {y} Td ({text}) Tj ET\n" for text, y in lines
    ).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(output)
    output += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    output += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    output += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    path.write_bytes(bytes(output))


def test_pdf_text_layer_positions_runs_on_the_page(tmp_path):
    from app.services.document_preview.pdf import pdf_text_layers

    source = tmp_path / "text.pdf"
    _text_pdf(source, [("Quarterly review", 700)])

    [layer] = pdf_text_layers(source, max_pages=1)

    assert layer is not None
    assert layer.width == pytest.approx(612)
    [run] = [run for run in layer.runs if "Quarterly" in run.text]
    assert run.left == pytest.approx(100 * 72 / 612, abs=1.0)
    # Baseline 700pt from the bottom of a 792pt page: the run sits near the top.
    assert 7 < run.top < 14
    assert run.font_size > 0


def test_office_preview_uses_the_installed_runtime(monkeypatch, tmp_path):
    from pptx import Presentation

    from app.services.document_preview import service as preview

    source = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Launch readiness"
    slide.notes_slide.notes_text_frame.text = "Open with the timeline."
    presentation.save(source)
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setattr(
        preview,
        "installed_runtime_for_preview",
        lambda: installer.InstalledRuntime("26.8.0", tmp_path / "soffice", tmp_path),
    )

    def fake_convert(document, workdir, **_kwargs):
        pdf = workdir / "document.pdf"
        _text_pdf(pdf, [("Launch readiness", 700)])
        return pdf

    monkeypatch.setattr(convert, "convert_to_pdf", fake_convert)

    rendered = preview.render_document_preview(source).read_text(encoding="utf-8")

    assert 'data-preview-renderer="libreoffice-26.8.0"' in rendered
    assert 'data-preview-label="Slide 1 — Launch readiness"' in rendered
    assert 'data-preview-notes="Open with the timeline."' in rendered
    assert ">Launch readiness</span>" in rendered


def test_office_preview_falls_back_when_conversion_fails(monkeypatch, tmp_path):
    from docx import Document

    from app.services.document_preview import service as preview

    source = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Fallback body")
    document.save(source)
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setattr(
        preview,
        "installed_runtime_for_preview",
        lambda: installer.InstalledRuntime("26.8.0", tmp_path / "soffice", tmp_path),
    )

    def failing_convert(*_args, **_kwargs):
        raise convert.ConversionError("boom")

    monkeypatch.setattr(convert, "convert_to_pdf", failing_convert)

    rendered = preview.render_document_preview(source).read_text(encoding="utf-8")

    assert "data-preview-renderer" not in rendered
    assert "Fallback body" in rendered


def test_download_resumes_a_partial_file_and_hashes_every_byte(runtime_env, tmp_path):
    import hashlib

    archive = tmp_path / "runtime.tar.gz"
    _archive(archive, {EXECUTABLE: b"MZ" * 50_000, "soffice/share/a.ttf": b"f" * 9_000})
    record = _record(archive)
    runtime_env(record)
    asset = manifest.current_asset()
    assert asset is not None
    partial = installer.partial_download_path(asset)
    partial.parent.mkdir(parents=True)
    payload = archive.read_bytes()
    partial.write_bytes(payload[: len(payload) // 2])
    installer._set_job(
        installer.InstallJob("downloading", "26.8.0", 0, len(payload), "now")
    )

    digest = asyncio.run(installer._download(asset, partial))

    assert digest == hashlib.sha256(payload).hexdigest()
    assert partial.read_bytes() == payload
    assert installer.runtime_status().job.bytes_done == len(payload)


def test_failed_download_keeps_bytes_and_cancel_discards_them(runtime_env, tmp_path):
    archive = tmp_path / "runtime.tar.gz"
    _archive(archive, {EXECUTABLE: b"MZ"})
    record = _record(archive)
    runtime_env(record)
    asset = manifest.current_asset()
    assert asset is not None
    partial = installer.partial_download_path(asset)

    async def interrupted(asset_, destination):
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"half")
        raise installer.RuntimeInstallError("connection reset")

    async def cancelled(asset_, destination):
        destination.write_bytes(b"half")
        raise asyncio.CancelledError

    original = installer._download
    try:
        installer._download = interrupted
        with pytest.raises(installer.RuntimeInstallError):
            asyncio.run(installer._install(asset))
        assert partial.read_bytes() == b"half"

        installer._download = cancelled
        with pytest.raises(asyncio.CancelledError):
            asyncio.run(installer._install(asset))
        assert not partial.exists()
    finally:
        installer._download = original


def test_stale_leftovers_are_removed_except_the_resumable_download(
    runtime_env, tmp_path
):
    archive = tmp_path / "runtime.tar.gz"
    _archive(archive, {EXECUTABLE: b"MZ"})
    runtime_env(_record(archive))
    asset = manifest.current_asset()
    assert asset is not None
    home = installer.runtime_home()
    (home / ".staging-old" / "tree").mkdir(parents=True)
    old_partial = home / f".download-{'0' * 64}.part"
    old_partial.write_bytes(b"old")
    keep = installer.partial_download_path(asset)
    keep.write_bytes(b"resume me")

    installer._remove_stale_staging(keep=keep)

    assert not (home / ".staging-old").exists()
    assert not old_partial.exists()
    assert keep.read_bytes() == b"resume me"

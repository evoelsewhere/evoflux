from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

import pytest
from pptx.util import Inches

from app.services.document_preview import service as preview
from app.services.document_preview.live_deck import deck_is_live, read_deck_plan

SCRIPT = (
    Path(__file__).resolve().parents[2]
    / "app/agent/builtin_skills/pptx-official/scripts/deck_live.py"
)


@pytest.fixture
def deck_live():
    spec = importlib.util.spec_from_file_location("deck_live", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules["deck_live"] = module
    spec.loader.exec_module(module)
    yield module
    sys.modules.pop("deck_live", None)


SESSION = "session-a"


@pytest.fixture(autouse=True)
def building_session(monkeypatch):
    # The agent's shell names its session; ``init`` records it in the plan.
    monkeypatch.setenv("EVOFLUX_SESSION", SESSION)


def _render(deck: Path, session: str | None = SESSION) -> str:
    """The preview as ``session`` sees it while its turn runs (None: idle)."""
    return preview._render_document_preview(deck, live_session=session).read_text(
        encoding="utf-8"
    )


def _slide_statuses(rendered: str) -> list[str]:
    return re.findall(r'<article[^>]*data-slide-status="(\w+)"', rendered)


def _fresh_slides(rendered: str) -> int:
    return len(re.findall(r"<article[^>]*\bdata-slide-fresh\b", rendered))


def _add_slide(prs, text: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = text


def test_live_deck_previews_finished_slides_and_placeholders(
    deck_live, monkeypatch, tmp_path
):
    deck = tmp_path / "deck.pptx"
    deck_live.init(deck, 3)
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))

    plan = read_deck_plan(deck)
    assert plan is not None and plan.total == 3 and not plan.done
    assert plan.session == SESSION
    rendered = _render(deck)
    assert 'data-deck-live="true"' in rendered
    assert _slide_statuses(rendered) == ["building", "pending", "pending"]
    assert 'data-preview-label="Slide 2"' in rendered
    assert "Building slide 1 of 3" in rendered
    assert _fresh_slides(rendered) == 0

    live = deck_live.LiveDeck(deck)
    prs = live.open()
    _add_slide(prs, "Cover text")
    live.save(prs)

    rendered = _render(deck)
    assert _slide_statuses(rendered) == ["done", "building", "pending"]
    assert "Cover text" in rendered
    assert "Building slide 2 of 3" in rendered
    # The slide that just landed is the only one marked fresh.
    assert _fresh_slides(rendered) == 1
    assert 'class="done" style="width:33.33%"' in rendered

    for text in ("Why now text", "Roadmap text"):
        _add_slide(prs, text)
    live.finish(prs)

    rendered = _render(deck)
    assert not deck_is_live(deck, SESSION)
    assert "data-deck-live" not in rendered
    assert _slide_statuses(rendered) == []


def _slide_file(folder: Path, name: str, text: str) -> Path:
    path = folder / name
    path.write_text(
        "from pptx.util import Inches\n"
        "from theme import LABEL\n\n"
        "def build(prs):\n"
        "    slide = prs.slides.add_slide(prs.slide_layouts[6])\n"
        "    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))"
        f".text = LABEL + {text!r}\n",
        encoding="utf-8",
    )
    return path


def _texts(deck: Path) -> list[str]:
    from pptx import Presentation

    return [
        "".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        for slide in Presentation(str(deck)).slides
    ]


def test_add_builds_one_slide_per_command(deck_live, monkeypatch, tmp_path):
    deck = tmp_path / "deck.pptx"
    slides = tmp_path / "slides"
    slides.mkdir()
    (slides / "theme.py").write_text("LABEL = '> '\n", encoding="utf-8")
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))
    deck_live.init(deck, 3)

    assert deck_live.add(deck, _slide_file(slides, "01_cover.py", "Cover")) == 1
    rendered = _render(deck)
    assert _slide_statuses(rendered) == ["done", "building", "pending"]

    deck_live.add(deck, _slide_file(slides, "02_why.py", "Why now"))
    deck_live.add(deck, _slide_file(slides, "03_roadmap.py", "Roadmap"))
    # Rebuild slide 2 in place: order and count stay, the old slide is gone.
    deck_live.add(deck, _slide_file(slides, "02_why.py", "Why now v2"), replace=2)
    assert _texts(deck) == ["> Cover", "> Why now v2", "> Roadmap"]
    assert deck_is_live(deck, SESSION)

    deck_live.mark(deck, done=True)
    assert not deck_is_live(deck, SESSION)


def test_add_leaves_the_deck_alone_when_a_slide_file_fails(deck_live, tmp_path):
    deck = tmp_path / "deck.pptx"
    deck_live.init(deck, 1)
    before = deck.read_bytes()
    broken = tmp_path / "01_cover.py"
    broken.write_text("def build(prs):\n    pass\n", encoding="utf-8")

    with pytest.raises(deck_live.DeckLiveError, match="exactly one slide"):
        deck_live.add(deck, broken)
    with pytest.raises(deck_live.DeckLiveError, match="--replace 3"):
        deck_live.add(deck, broken, replace=3)
    assert deck.read_bytes() == before


def test_mark_restores_the_plan_after_a_generator_rewrites_the_deck(
    deck_live, tmp_path
):
    from pptx import Presentation

    deck = tmp_path / "deck.pptx"
    deck_live.init(deck, 2)
    rewritten = Presentation()
    _add_slide(rewritten, "One")
    rewritten.save(deck)  # e.g. PptxGenJS writing slides 1..k from scratch
    assert read_deck_plan(deck) is None

    deck_live.mark(deck)

    plan = read_deck_plan(deck)
    assert plan is not None and plan.total == 2 and not plan.done
    assert plan.session == SESSION
    deck_live.mark(deck, done=True)
    assert not deck_is_live(deck, SESSION)


def test_live_deck_renders_natively_even_with_the_exact_renderer(
    deck_live, monkeypatch, tmp_path
):
    from app.services.office_runtime import installer

    deck = tmp_path / "deck.pptx"
    deck_live.init(deck, 1)
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setattr(
        preview,
        "installed_runtime_for_preview",
        lambda: installer.InstalledRuntime("26.8.0", tmp_path / "soffice", tmp_path),
    )

    def fail(*_args, **_kwargs):
        raise AssertionError("a deck under construction must not use LibreOffice")

    monkeypatch.setattr(preview, "_render_with_office_runtime", fail)

    assert preview._renderer_identity(deck, live=True) == "native-live"
    assert 'data-deck-live="true"' in _render(deck)


def test_only_the_building_session_sees_the_deck_live_while_it_runs(
    deck_live, monkeypatch, tmp_path
):
    deck = tmp_path / "deck.pptx"
    deck_live.init(deck, 2)
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))
    live = deck_live.LiveDeck(deck)
    prs = live.open()
    _add_slide(prs, "Cover text")
    live.save(prs)

    assert _slide_statuses(_render(deck)) == ["done", "building"]
    # Another session editing the same file sees just the finished slide…
    other = _render(deck, "session-b")
    assert "data-deck-live" not in other and "Cover text" in other
    assert "slide-skeleton" not in other
    # …and so does the building session once its turn is over (abandoned build).
    assert "data-deck-live" not in _render(deck, None)
    assert not deck_is_live(deck, "session-b")


def _skeletons(rendered: str) -> list[str]:
    return re.findall(r'<section class="slide slide-skeleton".*?</section>', rendered)


def test_every_skeleton_looks_the_same_whatever_the_plan(deck_live, tmp_path):
    deck = tmp_path / "deck.pptx"
    deck_live.init(deck, 4)
    # A plan written with titles (an older script) shows none of them.
    older = dict(
        deck_live._read_plan(deck), titles=["Cover", "Roadmap", "Pricing", "Thanks"]
    )
    deck_live._write(deck, deck.read_bytes(), older)

    rendered = _render(deck)
    pending = [
        re.sub(r"Slide \d+", "Slide N", skeleton)
        for skeleton in _skeletons(rendered)[1:]
    ]
    assert len(pending) == 3 and len(set(pending)) == 1
    assert "Roadmap" not in rendered and "Pricing" not in rendered


def test_init_needs_a_slide_count(deck_live, tmp_path):
    with pytest.raises(deck_live.DeckLiveError, match="--slides"):
        deck_live.init(tmp_path / "deck.pptx", 0)


def _built_deck(deck_live, tmp_path: Path, monkeypatch) -> tuple[Path, Path]:
    """A finished two-slide deck built from slide files."""
    deck = tmp_path / "deck.pptx"
    slides = tmp_path / "slides"
    slides.mkdir()
    (slides / "theme.py").write_text("LABEL = '> '\n", encoding="utf-8")
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))
    deck_live.init(deck, 2)
    deck_live.add(deck, _slide_file(slides, "01_cover.py", "Cover"))
    deck_live.add(deck, _slide_file(slides, "02_why.py", "Why"))
    deck_live.mark(deck, done=True)
    return deck, slides


def test_fixing_a_finished_deck_never_brings_the_loading_state_back(
    deck_live, monkeypatch, tmp_path
):
    deck, slides = _built_deck(deck_live, tmp_path, monkeypatch)

    # QA fixes one slide in place, then a shared helper for all of them.
    deck_live.add(deck, _slide_file(slides, "02_why.py", "Why v2"), replace=2)
    assert _texts(deck) == ["> Cover", "> Why v2"]
    assert not deck_is_live(deck, SESSION)
    # A different length, so bytecode cached this same second is not reused.
    (slides / "theme.py").write_text("LABEL = '## '\n", encoding="utf-8")
    assert deck_live.rebuild(deck, slides) == 2
    assert _texts(deck) == ["## Cover", "## Why v2"]

    rendered = _render(deck)
    assert not deck_is_live(deck, SESSION)
    assert "data-deck-live" not in rendered and "slide-skeleton" not in rendered


def test_fixes_work_even_when_the_scratch_plan_is_gone(
    deck_live, monkeypatch, tmp_path
):
    deck, slides = _built_deck(deck_live, tmp_path, monkeypatch)
    deck_live._plan_path(deck).unlink()  # e.g. the temp folder was cleaned

    deck_live.add(deck, _slide_file(slides, "01_cover.py", "Cover v2"), replace=1)
    deck_live.mark(deck)

    plan = read_deck_plan(deck)
    assert plan is not None and plan.done
    assert _texts(deck) == ["> Cover v2", "> Why"]


def test_init_refuses_to_empty_a_deck_that_has_slides(deck_live, monkeypatch, tmp_path):
    deck, _slides = _built_deck(deck_live, tmp_path, monkeypatch)

    with pytest.raises(deck_live.DeckLiveError, match="--replace N"):
        deck_live.init(deck, 2)
    assert _texts(deck) == ["> Cover", "> Why"]

    deck_live.init(deck, 2, force=True)
    assert _texts(deck) == []
    assert deck_is_live(deck, SESSION)

from __future__ import annotations

import json
from pathlib import Path
from typing import cast

from openpyxl import load_workbook
import pytest

from app.agent.builtin_plugins.documents.engines import xlsx as pipeline
from app.agent.builtin_plugins.documents.rendering.runtime import file_sha256


def _new_project() -> dict[str, object]:
    return {
        "schema_version": 1,
        "title": "Plan",
        "mode": "new",
        "template_confirmed": False,
        "sheets": [
            {
                "name": "Inputs",
                "create_if_missing": True,
                "operations": [
                    {
                        "operation": "write_range",
                        "range": "A1:B2",
                        "values": [["Label", "Value"], ["Revenue", 10]],
                    }
                ],
            }
        ],
    }


def test_workbook_project_requires_template_lineage() -> None:
    raw = _new_project()
    raw["mode"] = "template"
    with pytest.raises(ValueError, match="source_sha256"):
        pipeline.WorkbookProject.model_validate(raw)


def test_workbook_catalog_exposes_openxml_operations() -> None:
    catalog = pipeline.workbook_catalog()
    assert catalog["workflow"] == "editable-openxml-xlsx"
    assert {"write_range", "add_chart", "autofit_columns"} <= set(catalog["operations"])


def test_autofit_requires_a_bounded_range() -> None:
    raw = _new_project()
    sheets = cast(list[dict[str, object]], raw["sheets"])
    sheets[0]["operations"] = [{"operation": "autofit_columns", "range": "A:B"}]

    with pytest.raises(ValueError, match="bounded A1 notation"):
        pipeline.WorkbookProject.model_validate(raw)


@pytest.mark.asyncio
async def test_compose_creates_editable_workbook_and_preview(tmp_path: Path) -> None:
    raw = _new_project()
    sheets = cast(list[dict[str, object]], raw["sheets"])
    operations = cast(list[dict[str, object]], sheets[0]["operations"])
    operations.extend(
        [
            {
                "operation": "write_range",
                "range": "C1:C2",
                "formulas": [["=B1"], ["=B2*2"]],
            },
            {
                "operation": "style_range",
                "range": "A1:C1",
                "format": {
                    "fill": "#0F766E",
                    "font": {"bold": True, "color": "#FFFFFF"},
                },
            },
            {
                "operation": "add_chart",
                "chart_type": "bar",
                "source_range": "A1:B2",
                "start_cell": "E2",
                "end_cell": "K12",
                "title": "Values",
                "has_legend": False,
            },
        ]
    )
    project_path = tmp_path / "project.json"
    project_path.write_text(json.dumps(raw), encoding="utf-8")
    output = tmp_path / "output.xlsx"

    result = await pipeline.compose_xlsx_project(
        project_path,
        None,
        output,
        workspace_root=tmp_path,
        work_dir=tmp_path / "work",
    )

    assert result.passed is True
    assert result.output == output
    assert len(result.previews) == 1
    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook["Inputs"]
        assert sheet["C2"].value == "=B2*2"
        assert len(sheet._charts) == 1  # noqa: SLF001
        assert sheet["A1"].fill.fgColor.rgb.endswith("0F766E")
    finally:
        workbook.close()


def test_template_validation_detects_changed_source(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    source.write_bytes(b"original")
    project = pipeline.WorkbookProject.model_validate(
        {
            "schema_version": 1,
            "title": "Edit",
            "mode": "template",
            "source_sha256": file_sha256(source),
            "template_confirmed": True,
            "sheets": [{"name": "Sheet1", "operations": []}],
        }
    )
    source.write_bytes(b"changed")
    with pytest.raises(ValueError, match="changed after inspection"):
        pipeline.validate_workbook_project(project, source)

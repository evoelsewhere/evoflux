from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

import pytest

from app.services.document_preview import service as preview
from app.services.document_preview.live_deck import deck_is_live, read_deck_plan

SCRIPT = (
    Path(__file__).resolve().parents[2]
    / "app/agent/builtin_skills/xlsx-official/scripts/workbook_live.py"
)
SESSION = "session-a"


@pytest.fixture
def workbook_live():
    spec = importlib.util.spec_from_file_location("workbook_live", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules["workbook_live"] = module
    spec.loader.exec_module(module)
    yield module
    sys.modules.pop("workbook_live", None)


@pytest.fixture(autouse=True)
def building_session(monkeypatch, tmp_path):
    monkeypatch.setenv("EVOFLUX_SESSION", SESSION)
    monkeypatch.setattr(preview.settings, "EVOFLUX_CACHE_DIR", str(tmp_path / "cache"))


def _render(book: Path, session: str | None = SESSION) -> str:
    return preview._render_document_preview(book, live_session=session).read_text(
        encoding="utf-8"
    )


def _statuses(rendered: str) -> list[str]:
    return re.findall(r'<section[^>]*data-slide-status="(\w+)"', rendered)


def _labels(rendered: str) -> list[str]:
    return re.findall(r'<section[^>]*data-preview-label="([^"]+)"', rendered)


def _sheet_file(folder: Path, name: str, title: str, value: str) -> Path:
    path = folder / name
    path.write_text(
        "from theme import LABEL\n\n"
        "def build(wb):\n"
        f"    ws = wb.create_sheet({title!r})\n"
        f"    ws['A1'] = LABEL + {value!r}\n",
        encoding="utf-8",
    )
    return path


@pytest.fixture
def sheets(tmp_path: Path) -> Path:
    folder = tmp_path / "sheets"
    folder.mkdir()
    (folder / "theme.py").write_text("LABEL = '> '\n", encoding="utf-8")
    return folder


def test_live_workbook_shows_added_sheets_then_skeletons(
    workbook_live, sheets, tmp_path
):
    book = tmp_path / "book.xlsx"
    workbook_live.init(book, 3)

    plan = read_deck_plan(book)
    assert plan is not None and plan.total == 3 and plan.added == 0 and not plan.done
    rendered = _render(book)
    # The stand-in sheet a new workbook needs never shows.
    assert 'data-deck-live="true"' in rendered and "Building…" not in _labels(rendered)
    assert _statuses(rendered) == ["building", "pending", "pending"]
    assert "Building sheet 1 of 3" in rendered

    workbook_live.add(book, _sheet_file(sheets, "01_inputs.py", "Inputs", "rates"))
    rendered = _render(book)
    assert _statuses(rendered) == ["done", "building", "pending"]
    assert _labels(rendered)[0] == "Inputs" and "&gt; rates" in rendered
    assert len(re.findall(r"<section[^>]*\bdata-slide-fresh\b", rendered)) == 1

    workbook_live.add(book, _sheet_file(sheets, "02_model.py", "Model", "sum"))
    workbook_live.add(book, _sheet_file(sheets, "03_notes.py", "Notes", "why"))
    workbook_live.finish(book)
    rendered = _render(book)
    assert not deck_is_live(book, SESSION)
    assert "data-deck-live" not in rendered and _labels(rendered) == [
        "Inputs",
        "Model",
        "Notes",
    ]


def test_fixing_a_finished_workbook_keeps_it_finished(workbook_live, sheets, tmp_path):
    from openpyxl import load_workbook

    book = tmp_path / "book.xlsx"
    workbook_live.init(book, 2)
    workbook_live.add(book, _sheet_file(sheets, "01_inputs.py", "Inputs", "rates"))
    workbook_live.add(book, _sheet_file(sheets, "02_model.py", "Model", "sum"))
    workbook_live.finish(book)

    workbook_live.add(
        book, _sheet_file(sheets, "01_inputs.py", "Inputs", "rates v2"), replace=1
    )
    workbook = load_workbook(book)
    assert workbook.sheetnames == ["Inputs", "Model"]
    assert workbook["Inputs"]["A1"].value == "> rates v2"
    (sheets / "theme.py").write_text("LABEL = '## '\n", encoding="utf-8")
    assert workbook_live.rebuild(book, sheets) == 2
    assert load_workbook(book)["Model"]["A1"].value == "## sum"

    assert not deck_is_live(book, SESSION)
    assert "sheet-skeleton" not in _render(book)
    with pytest.raises(workbook_live.WorkbookLiveError, match="--replace N"):
        workbook_live.init(book, 2)


def test_only_the_building_session_sees_the_workbook_live(
    workbook_live, sheets, tmp_path
):
    book = tmp_path / "book.xlsx"
    workbook_live.init(book, 2)
    workbook_live.add(book, _sheet_file(sheets, "01_inputs.py", "Inputs", "rates"))

    assert _statuses(_render(book)) == ["done", "building"]
    other = _render(book, "session-b")
    assert "data-deck-live" not in other and "sheet-skeleton" not in other


def test_a_sheet_file_must_add_exactly_one_sheet(workbook_live, tmp_path):
    book = tmp_path / "book.xlsx"
    workbook_live.init(book, 1)
    before = book.read_bytes()
    broken = tmp_path / "01_broken.py"
    broken.write_text("def build(wb):\n    pass\n", encoding="utf-8")

    with pytest.raises(workbook_live.WorkbookLiveError, match="exactly one sheet"):
        workbook_live.add(book, broken)
    assert book.read_bytes() == before
